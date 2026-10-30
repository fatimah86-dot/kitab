#!/usr/bin/env python3
"""Build the label-free final Al-Qorin DOCX from the audited 2-layer DOCX.

The transformation is intentionally subtractive only.  It removes the
document title, every unit/Layer heading, the three Layer-1 labels, and
all footnote/unclear-print paragraphs.  Arabic, Latin, Terjemahan,
Media, Langkah, Tujuan, images, fonts, and paragraph direction are
left byte-for-byte as paragraph objects from the source DOCX.
"""

from __future__ import annotations

import argparse
import hashlib
import sys
import zipfile
from collections import Counter
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_INPUT = ROOT / "al_qorin/AL_QORIN_123_BERSIH_2_LAYER.docx"
DEFAULT_OUTPUT = ROOT / "al_qorin/AL_QORIN_123_BERSIH_FINAL_TANPA_SAMPAH.docx"

DELETE_STYLES = {"Title", "Heading 1", "Heading 2", "Heading 3"}
EXACT_DELETE = {
    "SAMPUL",
    "LAYER 1 - KETERANGAN HALAMAN:",
    "LAYER 1",
    "LAYER 2",
    "PARAGRAF",
    "AL-QORIN WAR RIYADHAH",
    "1. Arab:",
    "2. Latin:",
    "3. Terjemahan:",
}
CONTAINS_DELETE = (
    "tatacara praktik",
    "footnote",
    "[footnote",
    "cetakan/bacaannya tidak jelas",
    "«مصامات»",
    "«المتربرين»",
    "«اخلائق»",
    "«تنبتت»",
    "«الطبايح»",
    "«واعم»",
)


def normalized(text: str) -> str:
    return " ".join(text.split())


def has_drawing(paragraph) -> bool:
    return bool(paragraph._p.xpath('.//*[local-name()="drawing"]'))


def deletion_reason(paragraph) -> str | None:
    # Never remove a paragraph carrying one of the precision crops.
    if has_drawing(paragraph):
        return None

    if paragraph.style.name in DELETE_STYLES:
        return f"style:{paragraph.style.name}"

    text = normalized(paragraph.text)
    if text in EXACT_DELETE:
        return "exact"

    lowered = paragraph.text.lower()
    for needle in CONTAINS_DELETE:
        if needle.lower() in lowered:
            return f"contains:{needle}"
    return None


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def media_hashes(path: Path) -> list[str]:
    with zipfile.ZipFile(path) as archive:
        names = [n for n in archive.namelist() if n.startswith("word/media/")]
        return [hashlib.sha256(archive.read(n)).hexdigest() for n in names]


def font_part_count(path: Path) -> int:
    with zipfile.ZipFile(path) as archive:
        return len([n for n in archive.namelist() if n.startswith("word/fonts/")])


def build(input_path: Path, output_path: Path) -> Counter:
    input_path = input_path.resolve()
    output_path = output_path.resolve()
    if input_path == output_path:
        raise ValueError("output must not overwrite its input")
    if output_path.name != "AL_QORIN_123_BERSIH_FINAL_TANPA_SAMPAH.docx":
        raise ValueError("output filename must be AL_QORIN_123_BERSIH_FINAL_TANPA_SAMPAH.docx")

    doc = Document(input_path)
    removed: Counter = Counter()

    # Iterate over a snapshot because paragraph XML elements are removed.
    for paragraph in list(doc.paragraphs):
        reason = deletion_reason(paragraph)
        if reason is not None:
            removed[reason] += 1
            element = paragraph._element
            element.getparent().remove(element)

    # The document body no longer carries the old display title; keep the
    # package metadata consistent with that label-free output.
    doc.core_properties.title = ""

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)

    # Acceptance checks: no structural headings or banned displayed text,
    # while the retained practice apparatus and all 26 crops survive.
    out = Document(output_path)
    styles = Counter(p.style.name for p in out.paragraphs)
    if any(style in styles for style in DELETE_STYLES):
        raise AssertionError(f"structural styles remain: {styles}")

    for paragraph in out.paragraphs:
        text = normalized(paragraph.text)
        if text in EXACT_DELETE:
            raise AssertionError(f"exact banned line remains: {text!r}")
        lowered = paragraph.text.lower()
        for needle in CONTAINS_DELETE:
            if needle.lower() in lowered:
                raise AssertionError(f"banned text remains: {needle!r}")
        if "layer 1" in lowered or "layer 2" in lowered:
            raise AssertionError("a Layer label remains")

    if len(out.inline_shapes) != 26:
        raise AssertionError(f"expected 26 inline images, got {len(out.inline_shapes)}")
    if media_hashes(output_path) != media_hashes(input_path):
        raise AssertionError("embedded media hashes/order changed")
    if font_part_count(output_path) != 4:
        raise AssertionError("embedded font parts were not preserved")

    media_count = sum(normalized(p.text) == "Media:" for p in out.paragraphs)
    step_count = sum(normalized(p.text).startswith("Langkah ke-") for p in out.paragraphs)
    tujuan_count = sum(normalized(p.text).startswith("Tujuan") for p in out.paragraphs)
    if media_count != 72 or tujuan_count != 72 or step_count < 270:
        raise AssertionError(
            f"practice labels changed unexpectedly: media={media_count} "
            f"steps={step_count} tujuan={tujuan_count}"
        )

    removed["output_paragraphs"] = len(out.paragraphs)
    removed["media_labels"] = media_count
    removed["step_labels"] = step_count
    removed["tujuan_labels"] = tujuan_count
    return removed


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args(argv)

    stats = build(args.input, args.output)
    print(f"input={args.input.resolve()}")
    print(f"output={args.output.resolve()}")
    print(f"removed={dict(stats)}")
    print(f"bytes={args.output.resolve().stat().st_size}")
    print(f"sha256={sha256(args.output.resolve())}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
