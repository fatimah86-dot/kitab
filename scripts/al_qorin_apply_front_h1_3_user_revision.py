#!/usr/bin/env python3
"""Restructure the front of AL_QORIN_123_FINAL_KELUHAN_SELESAI.docx.

User revision (2026-10-30): Halaman 1 = cover (restored, verbatim user
text), Halaman 2-3 replaced entirely with the user's supplied
-Arabic/Latin/Terjemahan. Their two pages cover the old units 1-3
(boundary at maqamati / al-qashidin), so numbering from Halaman 4
onward is unchanged. New headings: 1 COVER KITAB, 2 PUJIAN PERTAMA
DAN LANJUTAN PUJIAN, 3 SYAHADAT DAN ALASAN PENULISAN (kept).
"""

from __future__ import annotations

import hashlib
import re
import sys
import zipfile
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
TARGET = ROOT / "al_qorin/AL_QORIN_123_FINAL_KELUHAN_SELESAI.docx"
SOURCE_FOR_MEDIA = ROOT / "al_qorin/AL_QORIN_123_BERSIH_2_LAYER.docx"

HEAD1 = "=== Halaman 1 — COVER KITAB ==="
HEAD2 = "=== Halaman 2 — PUJIAN PERTAMA DAN LANJUTAN PUJIAN ==="
HEAD3 = "=== Halaman 3 — SYAHADAT DAN ALASAN PENULISAN ==="

AR_COVER = [
    "الْقَرِينُ",
    "وَالرِّيَاضَاتُ الرُّوحَانِيَّةُ",
    "تَأْلِيفُ الْأُسْتَاذِ الْكَبِيرِ",
    "مُحَمَّد سَامِي مُحَمَّد أَحْمَد",
    "عُضْوِ الِاتِّحَادِ الْعَالَمِيِّ لِلْفَلَكِيِّينَ الرُّوحَانِيِّينَ",
    "الْمَقَرُّ بَارِيس",
    "النَّاشِر",
    "دَارُ الْفِكْرِ وَالثَّقَافَةِ",
    "مُحَمَّد سَامِي مُحَمَّد أَحْمَد",
]
LAT_COVER = [
    "Al-Qarīnu",
    "war-Riyāḍātur-Rūḥāniyyatu",
    "Ta'līful-ustāżil-kabīri",
    "Muḥammad Sāmī Muḥammad Aḥmad",
    "‘Uḍwil-ittiḥādil-‘ālamiyyi lil-falakiyyīnar-rūḥāniyyīna",
    "Al-Maqarru Bārīs",
    "An-Nāsyir",
    "Dārul-Fikri wats-Tsaqāfati",
    "Muḥammad Sāmī Muḥammad Aḥmad",
]
TERJ_COVER = [
    "QARIN",
    "DAN RIYADHAH-RIYADHAH SPIRITUAL",
    "Karya Al-Ustadz Al-Kabir (Maha Guru)",
    "Muhammad Sami Muhammad Ahmad",
    "Anggota Persatuan Ahli Astronomi Spiritual Internasional",
    "Markas: Paris",
    "Penerbit:",
    "Darul Fikr wa Ats-Tsaqafah (Rumah Pemikiran dan Kebudayaan)",
    "Muhammad Sami Muhammad Ahmad",
]

AR2 = [
    "بِسْمِ اللَّهِ الرَّحْمَنِ الرَّحِيمِ",
    "مُقَدِّمَةٌ :",
    "الْحَمْدُ لِلَّهِ الَّذِي رَفَعَ حُجُبَ أَسْتَارِ الْأَسْرَارِ عَنْ حَقَائِقِ بَصَائِرِ الْمُقَرَّبِينَ وَنَوَّرَ بِنُورِ الْهِدَايَةِ الْإِيمَانِيَّةِ مِشْكَاةَ قُلُوبِ الْعَارِفِينَ وَشَرَحَ بِسِرِّ الْمَوَاهِبِ الرَّبَّانِيَّةِ بَوَاطِنَ صُدُورِ الْمُحَقِّقِينَ وَأَسْكَرَ مِنْ خَالِصِ صِرْفِ مَحَبَّتِهِ جَوَاهِرَ أَرْوَاحِ الْمُحِبِّينَ وَهَدَى إِلَى أَقْرَبِ وَسِيلَةٍ سُلُوكَ مَقَامَاتِ الْمُخْلِصِينَ وَثَبَّتَ عَلَى صِرَاطِ الْحَقِيقَةِ أَقْدَامَ السَّالِكِينَ وَرَقَّى عَلَى مَنَابِرِ الْحِكَمِ النُّورَانِيَّةِ ذَوَاتِ أَفْكَارِ الْمُهْتَدِينَ وَنَزَّهَ فِي حَدِيقَةِ رِيَاضِ حَضْرَةِ قُدْسِهِ أَنْسَةَ أَلْبَابِ عُقُولِ الْمُتَبَرِّئِينَ وَتَقَرَّبَ بِمَحْضِ أَلْطَافِ عِنَايَةِ رَحْمَتِهِ بِخُضُوعِ خُشُوعِ الْمُنَاجِينَ .",
    "وَهُوَ الَّذِي تَفَرَّدَ بِسِرِّ الْوَحْدَانِيَّةِ وَتَوَحَّدَ بِتَحْقِيقِ الْحَقِيقَةِ الْأَحَدِيَّةِ وَاتَّصَفَ بِعِزَّةِ اللَّهِ الصَّمَدِيَّةِ وَتَنَزَّهَ عَنْ شِبْهِ الْمُلْحِدِينَ .",
    "الَّذِي مَنَّ بِرَوَاقِ التَّكْوِينِ عَلَى اخْتِلَافِ أَطْوَارِ مَقَامَاتِ",
]
LAT2 = [
    "Bismillāhir-raḥmānir-raḥīm. Muqaddimah:",
    "Al-ḥamdu lillāhil-lażī rafa‘a ḥujuba astāril-asrāri ‘an ḥaqā'iqi baṣā'iril-muqarrabīna wa nawwara binūril-hidāyatil-īmāniyyati misykāta qulūbil-‘ārifīna wa syaraḥa bisirril-mawāhibir-rabbāniyyati bawāṭina ṣudūril-muḥaqqiqīna wa askara min khāliṣi ṣirfi maḥabbatihī jawāhira arwāḥil-muḥibbīna wa hadā ilā aqrabi wasīlatin sulūka maqāmātil-mukhliṣīna wa tsabbata ‘alā ṣirāṭil-ḥaqīqati aqdāmas-sālikīna wa raqqā ‘alā manābiril-ḥikamin-nūrāniyyati żawāti afkāril-muhtadīna wa nazzaha fī ḥadīqati riyāḍi ḥaḍrati qudsihī ansata albābi ‘uqūlil-mutabarri'īna wa taqarraba bimaḥḍi alṭāfi ‘ināyati raḥmatihī bikhuḍū‘i khusyū‘il-munājīn.",
    "Wa huwal-lażī tafarrada bisirril-waḥdāniyyati wa tawaḥḥada bitaḥqīqil-ḥaqīqatil-aḥadiyyati wat-taṣafa bi‘izzatillāhiṣ-ṣamadiyyati wa tanazzaha ‘an syibhil-mulḥidīn.",
    "Al-Lażī manna birawāqit-takwīni ‘alā ikhtilāfi aṭwāri maqāmāti...",
]
TERJ2 = [
    "Dengan menyebut nama Allah Yang Maha Pengasih lagi Maha Penyayang.",
    "Pendahuluan:",
    "Segala puji bagi Allah yang telah menyingkap tabir-tabir rahasia dari hakikat mata batin orang-orang yang didekatkan (al-muqarrabin); yang menerangi relung hati para 'arifin (orang-orang yang makrifat) dengan cahaya petunjuk keimanan; yang melapangkan relung dada para muhaqqiqin dengan rahasia karunia ketuhanan; yang memabukkan permata ruh para pecinta dari kemurnian cinta-Nya yang murni; yang memberi petunjuk menuju wasilah terdekat untuk menempuh maqam-maqam orang-orang yang ikhlas; yang mengokohkan kaki para penempuh jalan (salikin) di atas titian hakikat; yang menaikkan esensi pikiran orang-orang yang mendapat petunjuk ke atas mimbar-mimbar hikmah yang bercahaya; yang menyucikan kedamaian inti akal orang-orang yang melepaskan diri (dari selain-Nya) di dalam taman riyadhah hadirat kesucian-Nya; dan yang mendekat dengan murni kelembutan inayah rahmat-Nya melalui ketundukan dan kekhusyukan orang-orang yang bermunajat.",
    "Dan Dialah Yang Maha Esa dengan rahasia keesaan (wahdaniyyah) dan menyendiri dengan perwujudan hakikat ke-Ahad-an (ahadiyyah), yang bersifat dengan keperkasaan Allah yang abadi (shamdiyyah), dan Maha Suci dari penyerupaan orang-orang yang menyimpang (mulhidin).",
    "Yang telah menganugerahkan hamparan penciptaan atas perbedaan tingkatan maqam...",
]

AR3 = [
    "الْقَاصِدِينَ وَبَسَطَ سَاحَةَ التَّحْكِيمِ بِصَفَاءِ جَوَاهِرِ أَنْفَاسِ أَرْوَاحِ الْمُتَمَكِّنِينَ الْقَدِيرُ الَّذِي وَسَمَ سُلْطَانَ الْعَقْلِ بِسِمَةِ الْعَجْزِ وَأَخْرَسَ فَصِيحَاتِ الْأَلْسُنِ بِيَدِ الْقَهْرِ أَحْمَدُهُ حَمْدَ الْمُقَرَّبِينَ وَأُؤْمِنُ إِيمَانَ الْعَارِفِينَ وَأَهْتَدِي بِهِدَايَةِ الْمُرِيدِينَ وَأَشْهَدُ أَنْ لَا إِلَهَ إِلَّا اللَّهُ وَحْدَهُ لَا شَرِيكَ لَهُ شَهَادَةَ مُنًى بِرِقِّ عُبُودِيَّتِي إِلَيْهِ .",
    "وَأَعْتَمِدُ فِي أَفْعَالِي وَأَقْوَالِي عَلَيْهِ وَأَشْهَدُ أَنَّ مُحَمَّداً عَبْدُهُ وَرَسُولُهُ الَّذِي اخْتَارَهُ مِنْ جَوَاهِرِ الْمُرْسَلِينَ وَخُلَاصَةِ أَنْوَارِ الصِّدِّيقِينَ وَحَقِيقَةِ حَقَائِقِ الشُّهَدَاءِ وَبَطَائِفِ أَسْرَارِ الصَّالِحِينَ صَلَّى اللَّهُ عَلَيْهِ حَامِلِي الْقُرْآنِ وَتَعَاقُبِ الْجَدِيدَانِ",
    "وَبَعْدُ فَإِنَّ جَمَاعَةً مِنَ الْمُحِبِّينَ الصَّادِقِينَ وَالْمُخْلِصِينَ الْخَالِصِينَ رَغِبُوا أَنْ أُبَيِّنَ لَهُمْ كَيْفِيَّةَ الرِّيَاضَاتِ وَتَرْبِيَتَ أَسْرَارِهَا فِي أُصُولِ الْمَوْجُودَاتِ وَمَآسِرِ الْبِدَايَاتِ مِنْهَا وَالنِّهَايَاتِ فَأَجَبْتُهُمْ إِلَى ذَلِكَ إِجَابَةَ مَنْ وَثِقَ بِحَلِّ مَوَدَّتِهِ فِي سِرِّ سِرِّهِ وَعَلَانِيَتِهِ وَنَجْوَاهُ بَعْدَ أَنِ اتَّخَذْتُ عُبُودِيَّةَ الذُّلِّ بَيَانَ الِاسْتِخَارَةِ .",
]
LAT3 = [
    "...al-qāṣidīna wa basaṭa sāḥatat-taḥkīmi biṣafā'i jawāhiri anfāsi arwāḥil-mutamakkinīnal-qadīrul-lażī wasama sulṭānal-‘aqli bisimatil-‘ajzi wa akhrasa faṣīḥātil-alsuni biyadil-qahri aḥmaduhū ḥamdal-muqarrabīna wa u'minu īmānal-‘ārifīna wa ahtadī bihidāyatil-murīdīna wa asyhadu al-lā ilāha illallāhu waḥdahū lā syarīka lahū syahādata munam-biriqqi ‘ubūdiyyatī ilaih.",
    "Wa a‘tamidu fī af‘ālī wa aqwālī ‘alaihi wa asyhadu anna muḥammadan ‘abduhū wa rasūluhul-lażī akhtārahū min jawāhiril-mursalīna wa khulāṣati anwāriṣ-ṣiddīqīna wa ḥaqīqati ḥaqā'iqisy-syuhadā'i wa laṭā'ifi asrāriṣ-ṣāliḥīna �ṣallallāhu ‘alaihi ḥāmilil-qur'āni wa ta‘āqubil-jadīdān.",
    "Wa ba‘du: Fa-inna jamā‘atam-minal-muḥibbīnaṣ-ṣādiqīna wal-mukhliṣīnal-khāliṣīna raghibū an ubayyina lahum kaifiyyatar-riyāḍāti wa tarbiyata asrārihā fī uṣūlil-maujūdāti wa ma'ātsiril-bidāyāti minhā wan-nihāyāti fa-ajabtuhum ilā żālika ijābata maw-watsiqa biḥalli mawaddatihī fī sirri sirrihī wa ‘alāniyatihī wa najwāhu ba‘da anit-takhażtu ‘ubūdiyyataz-zulli bayānal-istikhārah.",
]
TERJ3 = [
    "...orang-orang yang bermaksud (menuju Allah), dan membentangkan lapangan kekuasaan dengan kemurnian permata napas ruh orang-orang yang mapan kedudukannya (al-mutamakkinin). Yang Maha Kuasa, Yang telah menandai kekuasaan akal dengan tanda kelemahan, dan membisukan lisan-lisan yang fasih dengan Tangan (kekuasaan) Yang Menundukkan. Aku memuji-Nya dengan pujian orang-orang yang didekatkan, aku beriman dengan keimanan para 'arifin, aku mengambil petunjuk dengan petunjuk para murid (pencari jalan rohani), dan aku bersaksi bahwa tiada Tuhan selain Allah semata, tiada sekutu bagi-Nya, dengan persaksian pengharapan melalui ketertundukan penghambaanku kepada-Nya.",
    "Dan aku bersandar dalam perbuatan dan ucapanku kepada-Nya. Dan aku bersaksi bahwa Nabi Muhammad adalah hamba dan utusan-Nya yang Dia pilih dari permata para rasul, intisari cahaya para shiddiqin, hakikat dari hakikat para syuhada, dan kelembutan rahasia para shalihin. Semoga selawat Allah tercurah atas beliau (sepanjang masih ada) para pembawa Al-Qur'an dan bergantinya siang dan malam.",
    "Amma ba'du (Adapun setelah itu): Maka sesungguhnya sekelompok dari para pecinta yang jujur dan orang-orang ikhlas yang murni berharap agar aku menjelaskan kepada mereka tata cara riyadhah-riyadhah (latihan spiritual) dan mendidik rahasia-rahasianya di dalam akar wujud-wujud, beserta jejak-jejak (pengaruh) awalannya hingga akhirannya. Maka aku mengabulkan permintaan mereka tersebut dengan pengabulan seseorang yang percaya atas terurainya kasih sayang pada rahasia terdalamnya, keterbukaannya, dan bisikannya, setelah aku mengambil jalan penghambaan yang rendah diri melalui kejelasan istikharah.",
]


def has_drawing_el(el) -> bool:
    return bool(el.xpath('.//*[local-name()="drawing"]'))


def make_para(template_el, text: str):
    el = deepcopy(template_el)
    runs = el.findall(qn("w:r"))
    donor = deepcopy(runs[0]) if runs else None
    for child in list(el):
        if child.tag != qn("w:pPr"):
            el.remove(child)
    if donor is not None:
        texts = [t for t in donor.iter() if t.tag == qn("w:t")]
        for t in texts:
            t.text = ""
        if texts:
            texts[0].text = text
            texts[0].set(qn("xml:space"), "preserve")
            for extra in texts[1:]:
                extra.getparent().remove(extra)
        el.append(donor)
    return el


def media_hashes(path: Path) -> list[str]:
    with zipfile.ZipFile(path) as z:
        names = [n for n in z.namelist() if n.startswith("word/media/")]
        return [hashlib.sha256(z.read(n)).hexdigest() for n in names]


def main() -> int:
    doc = Document(TARGET)
    paras = list(doc.paragraphs)
    texts = [p.text for p in paras]
    heads = {}
    for i, t in enumerate(texts):
        m = re.match(r"^=== Halaman (\d+) — (.+) ===$", t.strip())
        if m:
            heads[int(m.group(1))] = i
    h1, h4 = heads[1], heads[4]
    if h1 != 0:
        raise AssertionError("document no longer starts at Halaman 1 heading")
    els = [p._element for p in paras]
    head4_el = els[h4]
    if any(has_drawing_el(els[i]) for i in range(h4)):
        raise AssertionError("front matter unexpectedly carries an image")

    head_tpl = deepcopy(els[0])
    arab_tpl = deepcopy(els[3])
    latin_tpl = deepcopy(els[6])
    terj_tpl = deepcopy(els[9])
    blank_tpl = deepcopy(els[10])
    if texts[10].strip() != "":
        raise AssertionError("blank template mismatch")

    body = head4_el.getparent()
    for i in range(h4):
        body.remove(els[i])

    sequence = []
    sequence.append(make_para(head_tpl, HEAD1))
    sequence += [make_para(arab_tpl, t) for t in AR_COVER]
    sequence += [make_para(latin_tpl, t) for t in LAT_COVER]
    sequence += [make_para(terj_tpl, t) for t in TERJ_COVER]
    sequence.append(make_para(blank_tpl, ""))
    sequence.append(make_para(head_tpl, HEAD2))
    sequence += [make_para(arab_tpl, t) for t in AR2]
    sequence += [make_para(latin_tpl, t) for t in LAT2]
    sequence += [make_para(terj_tpl, t) for t in TERJ2]
    sequence.append(make_para(blank_tpl, ""))
    sequence.append(make_para(head_tpl, HEAD3))
    sequence += [make_para(arab_tpl, t) for t in AR3]
    sequence += [make_para(latin_tpl, t) for t in LAT3]
    sequence += [make_para(terj_tpl, t) for t in TERJ3]
    sequence.append(make_para(blank_tpl, ""))
    for el in sequence:
        head4_el.addprevious(el)

    doc.save(TARGET)

    out = Document(TARGET)
    ot = [p.text for p in out.paragraphs]
    oh = {}
    for i, t in enumerate(ot):
        m = re.match(r"^=== Halaman (\d+) — (.+) ===$", t.strip())
        if m:
            oh[int(m.group(1))] = i
    if sorted(oh) != list(range(1, 123)):
        raise AssertionError("headings broken")
    if ot[oh[1]] != HEAD1 or ot[oh[2]] != HEAD2 or ot[oh[3]] != HEAD3:
        raise AssertionError("new headings mismatch")
    if ot[oh[1] + 1 : oh[2]] != AR_COVER + LAT_COVER + TERJ_COVER + [""]:
        raise AssertionError("cover unit mismatch")
    if ot[oh[2] + 1 : oh[3]] != AR2 + LAT2 + TERJ2 + [""]:
        raise AssertionError("unit2 mismatch")
    if ot[oh[3] + 1 : oh[4]] != AR3 + LAT3 + TERJ3 + [""]:
        raise AssertionError("unit3 mismatch")
    if len(out.inline_shapes) != 26:
        raise AssertionError("image count changed")
    if media_hashes(TARGET) != media_hashes(SOURCE_FOR_MEDIA):
        raise AssertionError("media hashes/order changed")

    print(f"total_paragraphs={len(ot)}")
    print(f"bytes={TARGET.stat().st_size}")
    print(f"sha256={hashlib.sha256(TARGET.read_bytes()).hexdigest()}")
    print("FRONT_H1_3_REVISION_QA_PASS")
    return 0


if __name__ == "__main__":
    sys.exit(main())
