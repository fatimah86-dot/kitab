#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Mechanical, logged restructure of Al-Qorin DRAF (5 old layers) into the locked
123-unit / 5-new-layer working master. Source sentences are reused; nothing is
rewritten here except headings, explicit-absence lines, crop markers, and the
removal classes mandated by STANDAR_MUTLAK_AL_QORIN_123.md.
Outputs: al_qorin/REVISI_AL_QORIN_123_KERJA.md, al_qorin/work/revised_units.json,
al_qorin/LOG_RESTRUKTUR_AWAL.md
"""
import json, re
from collections import Counter
from pathlib import Path
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
DRAF = ROOT / "google_drive/AL_QORIN_TEKS_KIRIMAN_SAMPUL_HAL122_DRAF (4).docx"
OUT_MD = ROOT / "al_qorin/REVISI_AL_QORIN_123_KERJA.md"
WORK = ROOT / "al_qorin/work"
LAYERS_OLD = ["1. ARAB", "2. LATIN", "3. TERJEMAHAN", "4. SYARAH DAN PENJELASAN", "5. PETUNJUK YANG ADA DALAM SUMBER"]
UNIT_START = lambda s: s == "SAMPUL" or bool(re.match(r"^M\d{1,3}\s+—", s))
LAYER_RE = re.compile(r"^([1-5])\.\s+(ARAB|LATIN|TERJEMAHAN|SYARAH DAN PENJELASAN|PETUNJUK YANG ADA DALAM SUMBER)\s*$")
NO_CALC = "«tidak ada hitungan di halaman ini»"

# ---------- paragraph markdown (preserve bold/italic runs) ----------
def para_md(p: Paragraph) -> str:
    out = ""
    for r in p.runs:
        t = r.text
        if not t:
            continue
        if r.bold or r.italic:
            core = t.strip()
            if core:
                lead = t[: len(t) - len(t.lstrip())]
                trail = t[len(t.rstrip()):]
                mark = "***" if (r.bold and r.italic) else ("**" if r.bold else "*")
                t = f"{lead}{mark}{core}{mark}{trail}"
        out += t
    # fall back if hyperlink/tab content was missed by runs
    plain_from_md = re.sub(r"\*+", "", out)
    if " ".join(plain_from_md.split()) != " ".join((p.text or "").split()):
        return p.text or ""
    return out

def cell_md(cell) -> str:
    parts = [para_md(p).strip().replace("|", "\\|") for p in cell.paragraphs]
    return "<br>".join([x for x in parts if x])

# ---------- body-order parse ----------
def parse():
    doc = Document(str(DRAF))
    pmap = {id(p._p): p for p in doc.paragraphs}
    header, units, cur, layer = [], [], None, None
    for child in doc.element.body:
        if child.tag == qn("w:p"):
            p = pmap.get(id(child)) or Paragraph(child, doc)
            plain = p.text or ""
            s = plain.strip()
            if UNIT_START(s):
                if cur: units.append(cur)
                cur = {"unit_id": "SAMPUL" if s == "SAMPUL" else s.split()[0],
                       "heading": s, "raw": {k: [] for k in LAYERS_OLD}}
                layer = None; continue
            lm = LAYER_RE.match(s)
            if cur is not None and lm:
                layer = f"{lm.group(1)}. {lm.group(2)}"; continue
            item = {"type": "p", "plain": plain, "md": para_md(p)}
            if cur is None: header.append(item)
            elif layer: cur["raw"][layer].append(item)
        elif child.tag == qn("w:tbl"):
            tb = Table(child, doc)
            rows = [[cell_md(c) for c in r.cells] for r in tb.rows]
            item = {"type": "table", "rows": rows}
            if cur is not None and layer: cur["raw"][layer].append(item)
    if cur: units.append(cur)
    return header, units

# ---------- sentence helpers ----------
PROTECT = {"QS.": "QS\u0001", "hal.": "hal\u0001", "Hal.": "Hal\u0001", "No.": "No\u0001"}
def split_sentences(text: str):
    t = " ".join((text or "").split())
    for a, b in PROTECT.items(): t = t.replace(a, b)
    parts = re.findall(r"\S.*?(?:[.!?؟]+[\"»”']?(?=\s|$)|$)", t)
    out = []
    for x in parts:
        for a, b in PROTECT.items(): x = x.replace(b, a)
        x = x.strip()
        if x: out.append(x)
    return out

CUES = re.compile(
    r"belum (?:dapat|pasti|ditetapkan|diverifikasi|terbaca|ada)|tidak dapat dipastikan|tidak (?:diam-diam|mengarang|direka|digambar ulang)|"
    r"dipertahankan|usulan|dibaca kerja|bacaan kerja|pembacaan kerja|bacaan awal|ekstrak|pada gambar|dari gambar|gambar mendukung|pindaian|"
    r"diperbaiki editorial|bersifat editorial|dinormalisasi|bentuk asli|bukan verifikasi|belum diverifikasi|rekaan|problematis|kejanggalan|"
    r"ditashih|tashih|keyakinan (?:sedang|rendah)|tanpa nomor|tidak terbaca|meragukan|tidak dianjurkan|tidak disahkan|bukan pengesahan|"
    r"tidak menjamin|rujuk (?:pindaian|gambar)|lihat (?:pindaian|gambar)|belum saya pastikan|tidak saya pastikan", re.I)
LABEL_RE = re.compile(r"(Nuqthah(?:/i‘rab)?|I‘rab|Shorof|Siyaq|Hukum(?: baca)?|Catatan(?: penutup| tashih)?|Tashih)\s*:")

def cue_sentences(text):
    return [s for s in split_sentences(text) if CUES.search(s) or "[CATATAN PENERJEMAH" in s]

def trim_dictation_lead(text):
    """Drop the quoted Arab/Latin segment that opens old dictation paragraphs;
    only the Indonesian commentary after ') —' can yield editor notes."""
    if text.startswith(("[رَسْمٌ", "[رسم")):
        return ""
    if text.startswith(("([Gambar", "[Gambar")):
        j = text.find(".]")
        return text[j + 2:].strip() if j >= 0 else ""
    if re.match(r"^(Nuqthah|I‘rab|Shorof|Siyaq|Hukum|Catatan|Tashih)\b", text):
        return text
    m = re.match(r"^.{0,1200}?\)\s+—\s+", text)
    if m and re.search(r"[\u0600-\u06ff]", m.group(0)) and (text[0] in "([«" or re.match(r"[\u0600-\u06ff]", text)):
        return text[m.end():].strip()
    return text

def salvage_paragraph(plain):
    notes = []
    text = trim_dictation_lead(" ".join(plain.split()))
    if not text:
        return notes
    spans = list(LABEL_RE.finditer(text))
    if not spans:
        return cue_sentences(text)
    pre = text[: spans[0].start()]
    notes += cue_sentences(pre)
    for i, m in enumerate(spans):
        end = spans[i + 1].start() if i + 1 < len(spans) else len(text)
        label, body = m.group(1), text[m.end():end].strip()
        if label.startswith("Nuqthah") or label.startswith("Catatan") or label.startswith("Tashih"):
            if body: notes.append(body)
        else:
            notes += cue_sentences(body)
    return notes

# ---------- old L5 processing ----------
EXPLAIN_START = ("Transkripsi ini", "Wafaq berpetak", "Pencatatan bentuk", "Yang dapat dipastikan")
def process_old_l5(items, stats):
    metode, notes, rajah = [], [], []
    i = 0
    while i < len(items):
        it = items[i]
        if it["type"] == "p" and it["plain"].strip().startswith(("RAJAH HAL.", "WAFAQ HAL.")):
            block = [it]; i += 1
            while i < len(items) and (items[i]["type"] == "table" or
                  (items[i]["type"] == "p" and items[i]["plain"].strip().startswith("WAFAQ HAL."))):
                block.append(items[i]); i += 1
            if i < len(items) and items[i]["type"] == "p" and items[i]["plain"].strip().startswith(EXPLAIN_START):
                block.append(items[i]); i += 1
            rajah.append(block); continue
        if it["type"] == "table":
            rajah.append([it]); i += 1; continue
        plain = it["plain"].strip(); i += 1
        if not plain: continue
        if plain.startswith("Qarinah "):
            stats["q_standalone_l5"] += 1; continue
        sents = split_sentences(plain)
        kept, removed_q = [], False
        for s in sents:
            st = s.strip()
            if st.startswith("Qarinah:") or st.startswith("Qarinah qabliah"):
                stats["q_inline_sentences"] += 1; removed_q = True; continue
            if st.startswith(("Catatan tashih:", "Harakat dan tanda baca", "Rajah, wafaq, dan jadwal")):
                notes.append(st); continue
            kept.append(st)
        if kept:
            metode.append(" ".join(kept) if removed_q else it["md"].strip())
    return metode, notes, rajah

# ---------- khususiyah ----------
TIME_RE = re.compile(r"\b(jam|waktu|malam|siang|pagi|sore|petang|fajar|subuh|zuhur|dhuhur|asar|ashar|maghrib|isya|senin|selasa|rabu|kamis|jumat|sabtu|minggu|ahad|tanggal|selama|hari)\b|ساعة", re.I)
MEDIA_RE = re.compile(r"\b(kertas|tinta|air mawar|za['’]faran|misik|mangkuk|piring|gelas|kain|benang|tali|telur|tanah liat|daun|kulit|tulang|logam|emas|perak|tembaga|besi|timah|raksa|minyak|lilin|wadah|bejana|botol|kaca|air sumur|ditulis|menulis|tulisan)\b", re.I)
BUKHUR_RE = re.compile(r"bukhur|بخور|kemenyan|luban|lubān|\bdupa\b|gaharu|\boud\b|dibakhurkan|membakhurkan|bakhurkan", re.I)
def khususiyah(metode_paras, no_practice):
    absent = "tidak disebutkan dalam sumber unit ini"
    if no_practice:
        return {"waktu": absent, "media": absent, "bukhur": absent}
    sents = []
    for p in metode_paras:
        sents += split_sentences(re.sub(r"\*+", "", p))
    def pick(rx):
        seen, out = set(), []
        for s in sents:
            if rx.search(s) and s not in seen:
                seen.add(s); out.append(s)
            if len(out) >= 5: break
        return " ".join(out) if out else absent
    return {"waktu": pick(TIME_RE), "media": pick(MEDIA_RE), "bukhur": pick(BUKHUR_RE)}

def render_table(rows):
    if not rows: return ""
    n = max(len(r) for r in rows)
    rows = [r + [""] * (n - len(r)) for r in rows]
    head = "| " + " | ".join(rows[0]) + " |"
    sep = "| " + " | ".join([":---:"] * n) + " |"
    body = ["| " + " | ".join(r) + " |" for r in rows[1:]]
    return "\n".join([head, sep] + body)

def main():
    header, units = parse()
    assert len(units) == 123, len(units)
    cropmap = json.loads((WORK / "crop_manifest.json").read_text(encoding="utf-8"))["crops"]
    insert_pages = {m["page"]: m for m in cropmap if m["insert"]}
    stats = Counter(); out_units = []
    jam_re = re.compile(r"\bjam\b|ساعة|sā‘ah|saa'ah|sa'ah", re.I)
    arith_re = re.compile(r"(?:\d+|[\u0660-\u0669]+)\s*\+\s*(?:\d+|[\u0660-\u0669]+)|dijumlah|penjumlahan|dibagi\s*(?:7|tujuh)", re.I)
    for u in units:
        raw = u["raw"]
        def paras(key): return [it["md"].strip() for it in raw[key] if it["type"] == "p" and it["plain"].strip()]
        # editor notes from old L4
        notes = []
        for it in raw["4. SYARAH DAN PENJELASAN"]:
            if it["type"] != "p": continue
            s = it["plain"].strip()
            if not s: continue
            if s.startswith("Qarinah "):
                stats["q_standalone_l4"] += 1; continue
            notes += salvage_paragraph(s)
        metode, notes5, rajah_blocks = process_old_l5(raw["5. PETUNJUK YANG ADA DALAM SUMBER"], stats)
        notes += notes5
        # dedupe notes preserving order
        ded, seen = [], set()
        for n in notes:
            n = " ".join(re.sub(r"\*+", "", n).split())
            # orphan image-caption fragments are not editor notes (the crop itself is inserted)
            if (n.startswith("[") and "lihat" in n.lower()) or (n.startswith("Khaatamu") and "[gambar" in n):
                continue
            if n and n not in seen: seen.add(n); ded.append(n)
        notes = ded
        no_practice = bool(metode) and metode[0].lstrip("* ").startswith("Tidak ada petunjuk praktik")
        if no_practice: stats["units_no_practice"] += 1
        if not metode:
            metode = ["Tidak ada metode praktik yang disebutkan dalam sumber unit ini."]; stats["units_metode_fallback"] += 1
        # jam screening on source layers 1-3 + metode
        srctext = "\n".join([it["plain"] for k in LAYERS_OLD[:3] for it in raw[k] if it["type"] == "p"] + metode)
        has_jam, has_arith = bool(jam_re.search(srctext)), bool(arith_re.search(srctext))
        jam_flag = has_jam and has_arith
        if has_jam: stats["units_jam_term"] += 1
        if has_arith: stats["units_arith"] += 1
        if jam_flag: stats["units_jam_arith_flag"] += 1
        jam_lines = ([f"Perlu penelaahan hitungan jam dari sumber unit ini sebelum dinyatakan final."] if jam_flag else [NO_CALC])
        # printed page / crop
        mpage = re.search(r"Halaman\s+(\d+)", u["heading"], re.I)
        page = int(mpage.group(1)) if mpage else None
        crop = insert_pages.get(page) if page else None
        if crop: stats["crops_inserted"] += 1
        kh = khususiyah(metode, no_practice)
        out_units.append({
            "unit_id": u["unit_id"], "heading": u["heading"], "page": page,
            "arab": paras("1. ARAB"), "latin": paras("2. LATIN"), "terjemahan": paras("3. TERJEMAHAN"),
            "metode": metode, "crop_page": (crop["page"] if crop else None),
            "crop_caption": (crop["caption"] if crop else None),
            "rajah_blocks": [[({"type": "p", "md": b["md"].strip()} if b["type"] == "p" else {"type": "table", "rows": b["rows"]}) for b in blk] for blk in rajah_blocks],
            "jam": jam_lines, "jam_flag": jam_flag, "catatan_penyunting": notes, "khususiyah": kh,
            "no_practice": no_practice,
        })
        stats["notes_total"] += len(notes)
        if not notes: stats["units_no_notes"] += 1
    (WORK / "revised_units.json").write_text(json.dumps(out_units, ensure_ascii=False, indent=1), encoding="utf-8")
    # ----- render MD -----
    head_plain = [h["plain"].strip() for h in header if h["plain"].strip()]
    md = ["# AL-QORIN WAR RIYADHAH — REVISI 123 ARTIKEL (KERJA)", "",
          "> Status: draf kerja hasil restrukturisasi mekanis terlog dari draf sumber; "
          "penyelesaian semantik dikerjakan per batch. Ikatan aturan: `../STANDAR_MUTLAK_AL_QORIN_123.md`.", ""]
    if len(head_plain) >= 4:
        md += ["> Catatan sumber (apa adanya): " + head_plain[2], "> " + head_plain[3], ""]
    for u in out_units:
        md += [f"## {u['heading']}", "", "### 1. Arab", ""]
        md += [x + "\n" for x in u["arab"]]
        md += ["### 2. Latin", ""] + [x + "\n" for x in u["latin"]]
        md += ["### 3. Terjemahan", ""] + [x + "\n" for x in u["terjemahan"]]
        md += ["### 4. Metode, Penentuan Jam, dan Catatan Penyunting", "", "#### 1. Metode", ""]
        md += [x + "\n" for x in u["metode"]]
        if u["crop_page"]:
            md += [f"![Crop halaman {u['crop_page']}](wafaq_qorin/{u['crop_page']}.jpg)", "",
                   f"[Crop halaman {u['crop_page']} — {u['crop_caption']}]", ""]
        for blk in u["rajah_blocks"]:
            for b in blk:
                md += [b["md"] + "\n"] if b["type"] == "p" else [render_table(b["rows"]), ""]
        md += ["#### 2. Penentuan Jam", ""] + [x + "\n" for x in u["jam"]]
        md += ["#### Catatan penyunting", ""]
        md += [x + "\n" for x in u["catatan_penyunting"]] or ["Tidak ada catatan penyunting tambahan pada unit ini.\n"]
        k = u["khususiyah"]
        md += ["### 5. Khususiyah — Waktu : Media : Bukhur", "",
               f"Waktu: {k['waktu']} : Media: {k['media']} : Bukhur: {k['bukhur']}", ""]
    text = "\n".join(md)
    OUT_MD.write_text(text, encoding="utf-8")
    # ----- log -----
    qlines = [l for l in text.splitlines() if l.strip().startswith("Qarinah ")]
    log = ["# LOG RESTRUKTUR AWAL AL-QORIN (mekanis, terlog)", "",
           f"- Unit: {len(out_units)}; lapis per unit: 5 (baru).",
           f"- Paragraf mandiri Qarinah dihapus: L4 lama {stats['q_standalone_l4']} + L5 lama {stats['q_standalone_l5']} = {stats['q_standalone_l4']+stats['q_standalone_l5']}.",
           f"- Kalimat Qarinah sebaris yang dibuang dari paragraf campuran: {stats['q_inline_sentences']}.",
           f"- Baris hasil yang masih diawali 'Qarinah ': {len(qlines)}.",
           f"- Catatan penyunting terselamatkan (paragraf): {stats['notes_total']}; unit tanpa catatan: {stats['units_no_notes']}.",
           f"- Label [CATATAN PENERJEMAH pada hasil: {text.count('[CATATAN PENERJEMAH')} (sumber: 25).",
           f"- Unit 'tidak ada petunjuk praktik': {stats['units_no_practice']}; metode cadangan eksplisit: {stats['units_metode_fallback']}.",
           f"- Crop wafaq tersisip: {stats['crops_inserted']} (setelah blok Metode).",
           f"- Saringan jam: unit mengandung istilah jam {stats['units_jam_term']}; mengandung pola aritmetika {stats['units_arith']}; keduanya (perlu telaah): {stats['units_jam_arith_flag']}.",
           f"- Frasa {NO_CALC} muncul: {text.count(NO_CALC)}.",
           "", "Catatan: perbedaan hitung Qarinah terhadap angka perintah (288) dirinci pada audit; "
           "yang dihapus di sini adalah seluruh paragraf mandiri berlabel Qarinah qabliah/ba‘diah yang terukur pada draf (4).", ""]
    (ROOT / "al_qorin/LOG_RESTRUKTUR_AWAL.md").write_text("\n".join(log), encoding="utf-8")
    print("\n".join(log))

if __name__ == "__main__":
    main()
