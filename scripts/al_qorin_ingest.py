#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Ingest DRAF Al-Qorin DOCX -> unit JSON (lossless paragraph text) + inventories.

The legacy FINAL-media mapping helper below is retained only as audit history;
it is not run by default and is not used for crops. Authorized wafaq images are
the user-uploaded precision crops staged by al_qorin_stage_user_crops.py.
Read-only w.r.t. sources in google_drive/.
"""
import json, re, sys, zipfile
from pathlib import Path
from docx import Document
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
DRAF = ROOT / "google_drive/AL_QORIN_TEKS_KIRIMAN_SAMPUL_HAL122_DRAF (4).docx"
FINAL = ROOT / "google_drive/KITAB AL-QORIN - EDISI REVISI LENGKAP 5 LAPIS - FINAL KHATAM.docx"
WORK = ROOT / "al_qorin/work"
WORK.mkdir(parents=True, exist_ok=True)

UNIT_RE = re.compile(r"^(SAMPUL|M(\d{1,3}))\s*(?:—|-)?\s*(.*)$")
LAYER_RE = re.compile(r"^([1-5])\.\s+(ARAB|LATIN|TERJEMAHAN|SYARAH DAN PENJELASAN|PETUNJUK YANG ADA DALAM SUMBER)\s*$")

def ingest_draf():
    d = Document(str(DRAF))
    paras = [p.text for p in d.paragraphs]
    # header = paras before first unit heading
    units, header, cur, layer = [], [], None, None
    for i, t in enumerate(paras):
        s = t.strip()
        m = UNIT_RE.match(s) if s else None
        # unit heading heuristic: SAMPUL alone or Mxx at line start; avoid matching prose
        if s == "SAMPUL" or (m and m.group(2) and re.match(r"^M\d{1,3}\s+—", s)):
            if cur: units.append(cur)
            cur = {"unit_id": s.split()[0] if s != "SAMPUL" else "SAMPUL",
                   "heading": s, "start_para": i, "layers": {}, "order": []}
            layer = None
            continue
        lm = LAYER_RE.match(s)
        if cur is not None and lm:
            layer = f"{lm.group(1)}. {lm.group(2)}"
            cur["layers"][layer] = []
            cur["order"].append(layer)
            continue
        if cur is None:
            header.append(t)
        elif layer is not None:
            cur["layers"][layer].append(t)
        else:
            cur.setdefault("pre", []).append(t)
    if cur: units.append(cur)
    # tables inventory
    tables = []
    for ti, tb in enumerate(d.tables):
        rows = [[c.text for c in r.cells] for r in tb.rows]
        tables.append({"index": ti, "rows": rows})
    out = {"source": str(DRAF.name), "header_paras": header, "units": units, "tables": tables}
    (WORK / "draf_units.json").write_text(json.dumps(out, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"DRAF units={len(units)} header_paras={len([h for h in header if h.strip()])} tables={len(tables)}")
    bad = [u["unit_id"] for u in units if len(u["order"]) != 5]
    print("units with !=5 layers:", bad[:20], "count", len(bad))
    ids = [u["unit_id"] for u in units]
    print("first/last ids:", ids[:4], ids[-4:])
    for name in ["Qarinah qabliah", "ba‘diah", "CATATAN PENERJEMAH", "Shorof:", "Hukum baca:", "penjumlahan"]:
        n = sum(p.count(name) for p in paras)
        print(f"count {name!r} =", n)
    qparas = [p for p in paras if p.strip().startswith("Qarinah")]
    print("paragraphs STARTING with 'Qarinah':", len(qparas))
    for q in qparas[:3]: print("  Q>", q[:140].replace("\n", " "))
    # where are the 4 tables? find body order context
    body = d.element.body
    seq = []
    for child in body:
        if child.tag == qn("w:tbl"):
            seq.append("[TABLE]")
        elif child.tag == qn("w:p"):
            txt = "".join(n.text or "" for n in child.iter() if n.tag == qn("w:t"))
            if txt.strip(): seq.append(txt.strip()[:60])
    for i, s in enumerate(seq):
        if s == "[TABLE]":
            print("TABLE context:", seq[max(0,i-2):i+2])
    for tb in tables:
        flat = " | ".join(c for r in tb["rows"] for c in r)[:300]
        print(f"table{tb['index']} {len(tb['rows'])}x{max((len(r) for r in tb['rows']), default=0)}: {flat}")
    return units

def map_final_media():
    d = Document(str(FINAL))
    rels = d.part.rels
    body = d.element.body
    events = []  # (kind, payload)
    cur_page = None
    last_texts = []
    for p in d.paragraphs:
        txt = p.text.strip()
        if p.style.name == "Halaman Kitab" or re.match(r"^(Halaman cetak \d+|Sampul —|Penutup —)", txt):
            cur_page = txt
        # drawings in this paragraph?
        blips = p._p.findall(".//" + qn("a:blip"))
        # extent
        exts = p._p.findall(".//" + qn("wp:extent"))
        for bi, b in enumerate(blips):
            rid = b.get(qn("r:embed"))
            if rid and rid in rels:
                target = rels[rid].target_ref
                ext = exts[bi] if bi < len(exts) else None
                cx = int(ext.get("cx")) if ext is not None else None
                cy = int(ext.get("cy")) if ext is not None else None
                events.append({"page_heading": cur_page, "target": target,
                               "para_text": txt[:120], "cx_emu": cx, "cy_emu": cy,
                               "prev_texts": last_texts[-3:]})
        if txt: last_texts.append(txt[:120])
    # media sizes/dims
    z = zipfile.ZipFile(str(FINAL))
    from PIL import Image
    import io
    info = {}
    for n in z.namelist():
        if n.startswith("word/media/"):
            data = z.read(n)
            try:
                im = Image.open(io.BytesIO(data)); dims = list(im.size)
            except Exception:
                dims = None
            info[n.split("/")[-1]] = {"bytes": len(data), "dims": dims}
    for e in events:
        base = e["target"].split("/")[-1]
        e["file"] = base; e.update(info.get(base, {}))
    (WORK / "final_media_map.json").write_text(json.dumps(events, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"FINAL image events={len(events)} unique files={len({e['file'] for e in events})} media total={len(info)}")
    unmapped = set(info) - {e["file"] for e in events}
    print("media not referenced in body paras:", sorted(unmapped)[:10], "count", len(unmapped))
    for e in events[:8]:
        print(" IMG>", e["file"], e.get("dims"), e["page_heading"], "| prev:", (e["prev_texts"][-1] if e["prev_texts"] else "")[:70])
    pindai = [e for e in events if any("Pindai sumber" in t for t in e["prev_texts"])]
    print("events preceded by 'Pindai sumber':", len(pindai))
    return events

if __name__ == "__main__":
    ingest_draf()
