#!/usr/bin/env python3
"""Replace Halaman 57-58 in AL_QORIN_123_FINAL_KELUHAN_SELESAI.docx.

User-supplied verbatim revision (2026-10-30), wafaq preserved: the
existing Mukhammas image paragraph in unit 58 keeps its element and
position (after the Terjemahan). The user's "4. Tata Cara Praktik"
section follows the wafaq. HTML entity &amp; is rendered as &.
Note: the new caption reading is yumha wa yusqa (user's text),
replacing the older supplied pattern entirely.
"""

from __future__ import annotations

import hashlib
import re
import sys
import zipfile
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
TARGET = ROOT / "al_qorin/AL_QORIN_123_FINAL_KELUHAN_SELESAI.docx"
SOURCE_FOR_MEDIA = ROOT / "al_qorin/AL_QORIN_123_BERSIH_2_LAYER.docx"

AR57 = [
    "الْفَصْلُ الثَّالِثَ عَشَرَ",
    "فَائِدَةٌ لِلتَّدَاوِي مِنْ جَمِيعِ الْأَمْرَاضِ",
    "بِسْمِ اللَّهِ الرَّحْمَنِ الرَّحِيمِ وَصَلَّى اللَّهُ عَلَى سَيِّدِنَا مُحَمَّدٍ وَعَلَى آلِهِ وَصَحْبِهِ وَسَلَّمَ ( لَقَدْ جَاءَكُمْ رَسُولٌ مِنْ أَنْفُسِكُمْ عَزِيزٌ عَلَيْهِ مَا عَنِتُّمْ ) إِلَى آخِرِ السُّورَةِ وَالْإِخْلَاصِ - وَالْمُعَوِّذَتَيْنِ - وَالْفَاتِحَةِ - ( وَنُنَزِّلُ مِنَ الْقُرْآنِ مَا هُوَ شِفَاءٌ وَرَحْمَةٌ لِلْمُؤْمِنِينَ ) - ( لَوْ أَنْزَلْنَا هَذَا الْقُرْآنَ عَلَى جَبَلٍ لَرَأَيْتَهُ ) إِلَى آخِرِ سُورَةِ الْحَشْرِ - اللَّهُمَّ أَنْتَ الشَّافِي وَأَنْتَ الْمُنْتَقِمُ - وَأَنْتَ الْمُعَافِي وَأَنْتَ الْمُجَلِّي - وَأَنْتَ الْوَاقِي وَأَنْتَ الْمُؤْلِمُ - وَأَنْتَ الرَّاقِي - خَلَقْتَنَا مِنْ سُلَالَةٍ مِنْ مَاءٍ مَهِينٍ - وَجَعَلْتَنَا فِي قَرَارٍ مَكِينٍ إِلَى قَدَرٍ مَعْلُومٍ - فَأَسْأَلُكَ بِأَسْمَائِكَ الْحُسْنَى وَذِكْرِكَ الْأَسْنَى وَصِفَاتِكَ الْعُلْيَا يَا مَنْ بِيَدِهِ الِابْتِلَاءُ وَالرُّقْيَا وَالْمَمَاتُ وَالْمَحْيَا وَالْعَافِيَةُ وَالشِّفَاءُ وَالْوَفَاءُ وَمَالِكَ جَمِيعِ الْأَشْيَاءِ وَخَالِقَهَا وَبَارِئَ جَمِيعِ الْأَحْيَاءِ وَرَازِقَهَا أَسْأَلُكَ بِحَقِّ حَبِيبِكَ الْأَعْظَمِ سَيِّدِنَا مُحَمَّدٍ صَلَّى اللَّهُ عَلَيْهِ وَسَلَّمَ وَبِمُعْجِزَاتِهِ وَبِحَقِّ خَلِيلِكَ إِبْرَاهِيمَ وَدَعْوَتِهِ وَبِحَقِّ رُوحِكَ عِيسَى وَرِفْعَتِهِ",
]
LAT57 = [
    "Al-Faṣluts-tsālitsa ‘asyara",
    "Fā'idatul-lit-tadāwī min jamī‘il-amrāḍ",
    "Bismillāhir-raḥmānir-raḥīm wa ṣallallāhu ‘alā sayyidinā muḥammadin wa ‘alā ālihī wa ṣaḥbihī wa sallam. (Laqad jā'akum rasūlum-min anfusikum ‘azīzun ‘alaihi mā ‘anittum) ilā ākhiris-sūrati, wal-ikhlāṣi, wal-mu‘awwiżataini, wal-fātiḥati. (Wa nunazzilu minal-qur'āni mā huwa syifā'uw-wa raḥmatul-lil-mu'minīna). (Lau anzalnā hāżal-qur'āna ‘alā jabalil-lara'aitahū) ilā ākhiri sūratil-ḥasyr.",
    "Allāhumma antasy-syāfī wa antal-muntaqimu, wa antal-mu‘āfī wa antal-mujallī, wa antal-wāqī wa antal-mu'limu, wa antar-rāqī. Khalaqtanā min sulālatim-mim-mā'im-mahīnin, wa ja‘altanā fī qarārim-makīnin, ilā qadarim-ma‘lūm.",
    "Fa-as'aluka bi'asmā'ikal-ḥusnā wa żikrikal-asnā wa ṣifātikal-‘ulyā, yā mam-biyadihil-ibtilā'u war-ruqyā wal-mamātu wal-maḥyā wal-‘āfiyatu wasy-syifā'u wal-wafā'u wa mālika jamī‘il-asyyā'i wa khāliQahā wa bāri'a jamī‘il-aḥyā'i wa rāziQahā. As'aluka biḥaqqi ḥabībikal-a‘ẓami sayyidinā muḥammadin ṣallallāhu ‘alaihi wa sallama wa bimu‘jizātihī, wa biḥaqqi khalīlika ibrāhīma wa da‘watihī, wa biḥaqqi rūḥika ‘īsā wa rif‘atihī...",
]
TERJ57 = [
    "Fasal Ketiga Belas",
    "Faidah Untuk Pengobatan Dari Segala Macam Penyakit",
    "Dengan menyebut nama Allah Yang Maha Pengasih lagi Maha Penyayang. Semoga selawat dan salam senantiasa tercurah kepada junjungan kita Nabi Muhammad, beserta keluarga dan para sahabatnya.",
    '(Bacalah ayat): "Sungguh telah datang kepadamu seorang Rasul dari kaummu sendiri, berat terasa olehnya penderitaanmu..." (QS. At-Taubah: 128) hingga akhir surah, lalu Surah Al-Ikhlas, Al-Mu\'awwidzatain (Al-Falaq & An-Nas), dan Al-Fatihah.',
    '(Lalu bacalah): "Dan Kami turunkan dari Al-Qur\'an suatu yang menjadi penawar/obat dan rahmat bagi orang-orang yang beriman." (QS. Al-Isra\': 82).',
    '(Lalu bacalah): "Kalau sekiranya Kami menurunkan Al-Qur\'an ini kepada sebuah gunung, pasti kamu akan melihatnya tunduk terpecah belah..." (QS. Al-Hasyr: 21) hingga akhir surah Al-Hasyr.',
    '"Ya Allah, Engkaulah Yang Maha Menyembuhkan dan Engkaulah Yang Maha Membalas. Engkaulah Yang Maha Menyelamatkan dan Engkaulah Yang Maha Menyingkap. Engkaulah Yang Maha Menjaga dan Engkaulah Yang Memberi rasa sakit. Engkaulah Yang Maha Mengobati (Meruqyah). Engkau telah menciptakan kami dari saripati air yang hina, lalu Engkau jadikan kami dalam tempat yang kokoh (rahim) sampai waktu yang ditentukan. Maka aku memohon kepada-Mu dengan Asma\'ul Husna-Mu, sebutan-Mu yang mulia, dan sifat-sifat-Mu yang luhur; Wahai Zat yang di Tangan-Nya (kekuasaan-Nya) segala cobaan, pengobatan, kematian, kehidupan, keselamatan, kesembuhan, dan pemenuhan janji. Wahai Pemilik segala sesuatu dan Penciptanya, wahai Yang Mengadakan segala makhluk hidup dan Yang Memberi rezeki kepada mereka. Aku memohon kepada-Mu demi hak kekasih-Mu yang agung, junjungan kita Nabi Muhammad SAW beserta mukjizat-mukjizatnya; demi hak kekasih-Mu Ibrahim dan seruannya; demi hak Ruh-Mu Isa dan derajatnya yang diangkat..."',
]

AR58 = [
    "وَبِحَقِّ صَفِيِّكَ نُوحٍ وَسَفِينَتِهِ وَمَا تَنَاسَلَ بَيْنَهُمَا مِنَ الْأَنْبِيَاءِ وَالْمُرْسَلِينَ اشْفِنِي بِشِفَائِكَ وَدَاوِنِي بِدَوَائِكَ وَعَافِنِي بِعَافِيَتِكَ وَبِمُعَافَاتِكَ فَإِنَّكَ لَا إِلَهَ إِلَّا أَنْتَ وَلَا رَبَّ سِوَاكَ وَالْحَوْقَلَةِ وَصَلَّى اللَّهُ عَلَى سَيِّدِنَا مُحَمَّدٍ وَعَلَى آلِهِ وَصَحْبِهِ .",
    '"مُخَمَّسٌ لِجَمِيعِ الْأَمْرَاضِ يُمْحَى وَيُسْقَى"',
]
LAT58 = [
    "...wa biḥaqqi ṣafiyyika nūḥin wa safīnatihī wa mā tanāsala bainahumā minal-ambiyā'i wal-mursalīna. Iṣyfinī bisyifā'ika, wa dāwinī bidawā'ika, wa ‘āfinī bi‘āfiyatika wa bimu‘āfātika, fa'innaka lā ilāha illā anta wa lā rabba siwāka. (Wal-ḥauqalatu). Wa ṣallallāhu ‘alā sayyidinā muḥammadin wa ‘alā ālihī wa ṣaḥbihī.",
    "Mukhammasun lijamī‘il-amrāḍi yumḥā wa yusqā.",
]
TERJ58 = [
    '"...dan demi hak hamba pilihan-Mu Nuh dan bahteranya, serta demi para nabi dan rasul yang silsilahnya lahir di antara keduanya. Sembuhkanlah aku dengan kesembuhan dari-Mu, obatilah aku dengan obat dari-Mu, dan selamatkanlah aku dengan keselamatan dan pemeliharaan-Mu. Karena sesungguhnya tiada tuhan melainkan Engkau, dan tiada tuhan (penguasa) selain-Mu." (Lalu bacalah Hauqalah: La haula wa la quwwata illa billahil \'aliyyil \'azhim). Dan semoga selawat Allah tercurah kepada junjungan kita Nabi Muhammad, beserta keluarga dan para sahabatnya.',
    "(Tabel Wafaq Mukhammas)",
    '"Ini adalah (wafaq) Mukhammas untuk seluruh penyakit, dihapus (dengan air) dan diminumkan."',
]

TATACARA = [
    "4. Tata Cara Praktik",
    "Berdasarkan Halaman 57 dan 58, fasal ini menjelaskan tata cara pengobatan menyeluruh (segala jenis penyakit medis maupun non-medis) melalui media air dan tulisan rajah (Wifiq).",
    "Langkah-Langkah Pengobatan:",
    "- Penulisan Wafaq & Ayat: Tulislah Wafaq Mukhammas (Tabel 5x5) yang tertera pada halaman 58 di atas media yang bersih (biasanya menggunakan piring putih polos atau kertas bersih dengan tinta za'faran).",
    '- Pengisian Kolom "At-Tawkil": Pada kotak tengah wafaq yang bertuliskan kata التوكيل (At-Tawkil), kolom tersebut tidak ditulis persis seperti kalimatnya, melainkan diisi dengan doa/perintah agar penyakit sang pasien diangkat (misal: Tawakkalu ya khuddama hazal wafaq bishifa\'i... [nama pasien]).',
    "- Pembacaan Doa: Setelah menulis wafaq, bacakan doa perlindungan dari Halaman 57 secara berurutan:",
    "  - Ayat Laqad Ja'akum (At-Taubah: 128-129).",
    "  - Surah Al-Ikhlas, Al-Falaq, An-Nas, dan Al-Fatihah.",
    "  - Ayat Wa nunazzilu minal qur'an... (Al-Isra': 82).",
    "  - Akhir Surah Al-Hasyr (ayat 21-24).",
    "  - Doa munajat kesembuhan (dari kalimat Allahumma Antash Syafi... hingga selesai di halaman 58).",
    "  - Ditutup dengan Hauqalah dan Selawat Nabi.",
    "- Peleburan (Al-Mahwu): Sesuai instruksi di bawah tabel wafaq (Yumha wa Yusqa), tulisan rajah tersebut kemudian dilebur/dilarutkan menggunakan air bersih.",
    "- Diminumkan: Air lunturan rajah tersebut kemudian diminumkan kepada orang yang sakit.",
]


def has_drawing_el(el) -> bool:
    return bool(el.xpath('.//*[local-name()="drawing"]'))


def make_para(template_el, text: str):
    el = deepcopy(template_el)
    runs = el.findall(qn("w:r"))
    donor = deepcopy(runs[0]) if runs else None
    for child in list(el):
        if child.tag != qn("w:pPr"):
            el.remove(child)
    if donor is not None:
        texts = [t for t in donor.iter() if t.tag == qn("w:t")]
        for t in texts:
            t.text = ""
        if texts:
            texts[0].text = text
            texts[0].set(qn("xml:space"), "preserve")
            for extra in texts[1:]:
                extra.getparent().remove(extra)
        el.append(donor)
    return el


def media_hashes(path: Path) -> list[str]:
    with zipfile.ZipFile(path) as z:
        names = [n for n in z.namelist() if n.startswith("word/media/")]
        return [hashlib.sha256(z.read(n)).hexdigest() for n in names]


def main() -> int:
    doc = Document(TARGET)
    paras = list(doc.paragraphs)
    texts = [p.text for p in paras]
    heads = {}
    for i, t in enumerate(texts):
        m = re.match(r"^=== Halaman (\d+) — (.+) ===$", t.strip())
        if m:
            heads[int(m.group(1))] = i
    h57, h58, h59 = heads[57], heads[58], heads[59]
    els = [p._element for p in paras]
    head57_el, head58_el, head59_el = els[h57], els[h58], els[h59]

    img_els = [els[i] for i in range(h57, h59) if has_drawing_el(els[i])]
    if len(img_els) != 1 or not (h58 < els.index(img_els[0]) < h59):
        raise AssertionError("expected exactly 1 wafaq image inside unit 58")
    img_el = img_els[0]

    arab_tpl = deepcopy(els[h57 + 1])
    latin_tpl = deepcopy(els[h57 + 5])
    terj_tpl = deepcopy(els[h57 + 7])
    if not texts[h57 + 5].startswith("Al-fashluts"):
        raise AssertionError("latin template mismatch")
    blank_tpl = deepcopy(els[h58 - 1])
    if texts[h58 - 1].strip() != "":
        raise AssertionError("blank template mismatch")

    body = head57_el.getparent()
    removed = 0
    node = head57_el.getnext()
    while node is not None and node is not head59_el:
        nxt = node.getnext()
        if node is not img_el and node is not head58_el and node.tag == qn("w:p"):
            body.remove(node)
            removed += 1
        node = nxt

    new57 = [make_para(arab_tpl, t) for t in AR57]
    new57 += [make_para(latin_tpl, t) for t in LAT57]
    new57 += [make_para(terj_tpl, t) for t in TERJ57]
    for el in new57:
        head58_el.addprevious(el)
    head58_el.addprevious(make_para(blank_tpl, ""))

    new58 = [make_para(arab_tpl, t) for t in AR58]
    new58 += [make_para(latin_tpl, t) for t in LAT58]
    new58 += [make_para(terj_tpl, t) for t in TERJ58]
    for el in new58:
        img_el.addprevious(el)

    anchor = img_el
    for t in TATACARA:
        el = make_para(terj_tpl, t)
        anchor.addnext(el)
        anchor = el
    anchor.addnext(make_para(blank_tpl, ""))

    doc.save(TARGET)

    out = Document(TARGET)
    ot = [p.text for p in out.paragraphs]
    oh = {}
    for i, t in enumerate(ot):
        m = re.match(r"^=== Halaman (\d+) — (.+) ===$", t.strip())
        if m:
            oh[int(m.group(1))] = i
    if sorted(oh) != list(range(1, 123)):
        raise AssertionError("headings broken")
    if ot[oh[57] + 1 : oh[58]] != AR57 + LAT57 + TERJ57 + [""]:
        raise AssertionError("unit57 mismatch")
    exp58 = AR58 + LAT58 + TERJ58 + [""] + TATACARA + [""]
    got58 = ot[oh[58] + 1 : oh[59]]
    if got58 != exp58:
        for a, b in zip(got58, exp58):
            if a != b:
                print("MISMATCH:\n got:", repr(a[:100]), "\n exp:", repr(b[:100]))
        raise AssertionError(f"unit58 mismatch len {len(got58)} vs {len(exp58)}")
    if len(out.inline_shapes) != 26:
        raise AssertionError("image count changed")
    if media_hashes(TARGET) != media_hashes(SOURCE_FOR_MEDIA):
        raise AssertionError("media hashes/order changed")

    print(f"removed_old_paras={removed}")
    print(f"total_paragraphs={len(ot)}")
    print(f"bytes={TARGET.stat().st_size}")
    print(f"sha256={hashlib.sha256(TARGET.read_bytes()).hexdigest()}")
    print("H57_58_REVISION_QA_PASS")
    return 0


if __name__ == "__main__":
    sys.exit(main())
