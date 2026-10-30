#!/usr/bin/env python3
"""Hapus tanda/kata yang didaftar pengguna dan terapkan dua penggantian.

Target hanya AL_QORIN_123_FINAL_KELUHAN_SELESAI.docx.

Penghapusan dilakukan pada bunyi persis dan varian ejaan yang jelas sama
pada daftar pengguna: penanda wafak beserta kurungnya, disclaimer
penerjemah, serta frasa kurung yang disebut satu per satu. Simbol/angka
Arab yang merupakan isi sumber tanpa kurung tidak dihapus. Paragraf yang
sesudah penghapusan hanya berisi penanda itu (atau tanda baca sisa)
dihapus sebagai paragraf agar dokumen rapih; paragraf gambar tidak
pernah dihapus.

Penggantian:
- `(Hayya Syarahiya, Adunai, Ashba-uuti, El Syaddai)` menjadi
  `(Ahyan Syarahiyan, Adunayya, Ashba-uuti, Ali Syaddāyya)`.
- setiap `Media:` menjadi `Pratik`.
"""

from __future__ import annotations

import hashlib
import re
import zipfile
from collections import Counter
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
TARGET = ROOT / "al_qorin/AL_QORIN_123_FINAL_KELUHAN_SELESAI.docx"

DISCLAIMER = "[CATATAN PENERJEMAH: Dicatat apa adanya, tidak dianjurkan, tidak disahkan sebagai amalan]"
OLD_NAMES = "(Hayya Syarahiya, Adunai, Ashba-uuti, El Syaddai)"
NEW_NAMES = "(Ahyan Syarahiyan, Adunayya, Ashba-uuti, Ali Syaddāyya)"

WAFAQ_PAREN_RE = re.compile(r"\(\s*\[\s*ada\s+wafa[kq]\s*\]\s*\)", re.I)
WAFAQ_SQUARE_RE = re.compile(r"\[\s*ada\s+wafa[kq]\s*\]", re.I)
WAFAQ_PLAIN_PAREN_RE = re.compile(r"\(\s*ada\s+wafa[kq]\s*\)", re.I)

# Urutan penting: bentuk kurung yang lebih panjang lebih dulu.
DELETE_PATTERNS: list[tuple[str, re.Pattern[str]]] = [
    ("frasa sebagian istilah", re.compile(r"\(sebagian istilah tidak pasti\)", re.I)),
    (
        "rangkaian simbol kurung",
        re.compile(r"\(✡\s*/\|\\\s*[مm]\s*#\s*١١١١\s*هـ\s*و\)"),
    ),
    ("simbol garis kurung", re.compile(r"\(\s*/\|\\\s*\)")),
    ("angka ١١١١ kurung", re.compile(r"\(١١١١\)")),
    ("penanda wafak dalam kurung", WAFAQ_PAREN_RE),
    ("penanda wafak kurung siku", WAFAQ_SQUARE_RE),
    ("penanda wafak kurung biasa", WAFAQ_PLAIN_PAREN_RE),
    (
        "angka 594/894 kurung",
        re.compile(r"\(5942131265\s+(?:dan|wa)\s+8941111\)", re.I),
    ),
    (
        "catatan nama surah",
        re.compile(r"\(nama surah terakhir pada terjemahan tidak pasti\)", re.I),
    ),
    ("lihat crop 121", re.compile(r"\(lihat crop Hal 121\)", re.I)),
    ("disclaimer penerjemah", re.compile(re.escape(DISCLAIMER))),
    ("gambar di Hal 109", re.compile(r"\(gambar di Hal 109\)", re.I)),
    ("lihat gambar Hal 109", re.compile(r"\(lihat gambar Hal 109\)", re.I)),
]

EXPECTED_DELETE_COUNTS = {
    "frasa sebagian istilah": 1,
    "rangkaian simbol kurung": 4,
    "simbol garis kurung": 1,
    "angka ١١١١ kurung": 1,
    # 20 berbentuk (...) dan 28 sisanya hanya [...]; total penanda 48.
    "penanda wafak dalam kurung": 20,
    "penanda wafak kurung siku": 28,
    "penanda wafak kurung biasa": 0,
    "angka 594/894 kurung": 2,
    "catatan nama surah": 1,
    "lihat crop 121": 1,
    "disclaimer penerjemah": 23,
    "gambar di Hal 109": 1,
    "lihat gambar Hal 109": 1,
}


def media_hashes(path: Path) -> list[str]:
    with zipfile.ZipFile(path) as z:
        return [
            hashlib.sha256(z.read(name)).hexdigest()
            for name in z.namelist()
            if name.startswith("word/media/")
        ]


def image_count(doc: Document) -> int:
    return sum(
        bool(p._p.xpath('.//*[local-name()="drawing"]')) for p in doc.paragraphs
    )


def has_drawing(paragraph) -> bool:
    return bool(paragraph._p.xpath('.//*[local-name()="drawing"]'))


def set_para_text(paragraph, text: str) -> None:
    nodes = paragraph._p.findall(".//" + qn("w:t"))
    if not nodes:
        paragraph.add_run(text)
        return
    nodes[0].text = text
    nodes[0].set(qn("xml:space"), "preserve")
    for node in nodes[1:]:
        node.text = ""


def clean_after_delete(text: str) -> str:
    # Rapikan hanya sisa mekanis dari penghapusan substring.
    text = re.sub(r"[ \t]{2,}", " ", text)
    text = re.sub(r"\s+([,.;:!?])", r"\1", text)
    text = re.sub(r"\(\s+\)", "", text)
    text = re.sub(r"\[\s*\]", "", text)
    text = re.sub(r"\s+([\)\]])", r"\1", text)
    text = re.sub(r"([\(\[])\s+", r"\1", text)
    return text.strip()


def transform(text: str, counts: Counter[str]) -> str:
    # Penggantian nama dilakukan persis pada bentuk kurung yang disebut.
    n = text.count(OLD_NAMES)
    if n:
        counts["ganti nama Hayya"] += n
        text = text.replace(OLD_NAMES, NEW_NAMES)

    deleted_here = False
    for label, pattern in DELETE_PATTERNS:
        text, n = pattern.subn("", text)
        if n:
            counts[label] += n
            deleted_here = True

    if deleted_here:
        text = clean_after_delete(text)

    # Penggantian label persis sebagaimana diminta pengguna.
    n = text.count("Media:")
    if n:
        counts["ganti Media:"] += n
        text = text.replace("Media:", "Pratik")
    return text


def main() -> int:
    before_media = media_hashes(TARGET)
    doc = Document(TARGET)
    before_paras = list(doc.paragraphs)
    before_texts = [p.text for p in before_paras]
    before_images = image_count(doc)
    assert before_images == 26, before_images

    counts: Counter[str] = Counter()
    intended: dict[int, str] = {}
    remove_indices: list[int] = []

    for idx, old in enumerate(before_texts):
        new = transform(old, counts)
        if new == old:
            continue
        # Paragraf yang isinya hanya penanda terdaftar (atau tinggal
        # tanda baca) dihapus utuh. Paragraf bergambar tidak disentuh.
        if not has_drawing(before_paras[idx]) and (
            not new.strip() or re.fullmatch(r"[.\s:;,-]+", new.strip())
        ):
            remove_indices.append(idx)
        else:
            intended[idx] = new

    for label, expected in EXPECTED_DELETE_COUNTS.items():
        assert counts[label] == expected, (label, counts[label], expected)
    assert counts["ganti nama Hayya"] == 1, counts
    assert counts["ganti Media:"] == 47, counts

    expected_after = [
        intended.get(i, text)
        for i, text in enumerate(before_texts)
        if i not in set(remove_indices)
    ]

    for idx, new in intended.items():
        set_para_text(before_paras[idx], new)
    for idx in reversed(remove_indices):
        el = before_paras[idx]._element
        el.getparent().remove(el)

    after_texts_in_memory = [p.text for p in doc.paragraphs]
    assert after_texts_in_memory == expected_after
    assert image_count(doc) == before_images

    doc.save(TARGET)

    check = Document(TARGET)
    check_texts = [p.text for p in check.paragraphs]
    assert check_texts == expected_after
    assert image_count(check) == before_images
    assert media_hashes(TARGET) == before_media

    joined = "\n".join(check_texts)
    for _label, pattern in DELETE_PATTERNS:
        assert not pattern.search(joined), pattern.pattern
    assert OLD_NAMES not in joined
    assert joined.count(NEW_NAMES) == 1
    assert "Media:" not in joined
    assert joined.count("Pratik") >= 47

    digest = hashlib.sha256(TARGET.read_bytes()).hexdigest()
    print(
        "REMOVE_LISTED_MARKERS_PASS "
        f"changed_paragraphs={len(intended)} removed_paragraphs={len(remove_indices)} "
        f"paragraphs={len(before_texts)}->{len(check_texts)} images={before_images} "
        f"bytes={TARGET.stat().st_size} sha256={digest}"
    )
    for label in list(EXPECTED_DELETE_COUNTS) + ["ganti nama Hayya", "ganti Media:"]:
        print(f"{label}: {counts[label]}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
