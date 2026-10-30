#!/usr/bin/env python3
"""Apply the user's Halaman 4-5 revision to the final KELUHAN DOCX.

User-supplied verbatim Arabic/Latin/Terjemahan (2026-10-30) replaces the
Layer-1 paragraphs of units 4 and 5 in
AL_QORIN_123_FINAL_KELUHAN_SELESAI.docx. Headings stay; the Halaman 4
Media/Langkah/Tujuan block is kept (it still matches the revised text:
starting with al-Fatihah for the Prophet, Ahl al-Bait, and the author).
The "[Halaman 4]/[Halaman 5]" markers in the user's message are not
content and are excluded.
"""

from __future__ import annotations

import hashlib
import re
import sys
import zipfile
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
TARGET = ROOT / "al_qorin/AL_QORIN_123_FINAL_KELUHAN_SELESAI.docx"
SOURCE_FOR_MEDIA = ROOT / "al_qorin/AL_QORIN_123_BERSIH_2_LAYER.docx"

AR4 = [
    "فَأَفَاضَ عَلَيَّ مِنْ أَنْوَارِ رَحْمَتِهِ قَبُولَ الْإِجَابَةِ وَكَشَفَ لِي عَنْ أَسْرَارِ أَبْوَابِ الْإِنَابَةِ عَلَى يَدِ الْعُلَمَاءِ الْمُتَقَدِّمِينَ وَالْأَئِمَّةِ وَالْمُهْتَدِينَ رَضِيَ اللَّهُ عَنْهُمْ ذَكَرُوا ذَلِكَ عَلَى أَنْوَاعٍ مِنَ الْأَمْثِلَةِ وَاخْتِلَافِ أَطْوَالِ الْأَبْنِيَةِ وَلَمْ يُودِعُوهُ مِنَ الْأَذْكَارِ أَنْوَاراً ثُمَّ أَنَّهُمْ ذَكَرُوا التَّدْرِيجَ عَلَى أَوْزَانِ الْقِسْطَاسِ وَلَمْ يُنَبِّهُوا عَلَى كَيْفِيَّةِ الِاحْتِرَازِ لِاخْتِلَالِ الْحَوَاسِّ .",
    "وَإِنَّنِي أُقَدِّمُ هَذِهِ الرِّيَاضَاتِ الرُّوحَانِيَّةَ الَّتِي لَيْسَ لَهَا مَثِيلٌ خِدْمَةً إِلَى الْمُؤْمِنِينَ فِي كُلِّ مَكَانٍ وَاللَّهُ الْمُوَفِّقُ وَأَرْجُو مِنَ اللَّهِ أَنْ يَتَقَبَّلَ هَذَا الْكِتَابَ وَمَا فِيهِ بِقَبُولٍ حَسَنٍ وَإِنِّي أَرْجُو مِنَ اللَّهِ أَنْ يَتَقَبَّلَنِي عِنْدَهُ وَيَتَقَبَّلَ هَذَا الْعَمَلَ مِنَ الْعَبْدِ الْفَقِيرِ إِلَى الْمَوْلَى الْجَلِيلِ وَمَنْ يُرِيدُ أَنْ يَصِلَ إِلَى الْمُرَادِ فَلْيَبْدَأْ بِقِرَاءَةِ الْفَاتِحَةِ لِسَيِّدِنَا مُحَمَّدٍ وَآلِ الْبَيْتِ وَالْمُؤَلِّفِ وَلِلَّهِ الْحَمْدُ فِي الْأُولَى وَالْآخِرَةِ .",
    "الْمُؤَلِّف",
    "مُحَمَّد سَامِي مُحَمَّد أَحْمَد",
]

LAT4 = [
    "Fa-afāḍa ‘alayya min anwāri raḥmatihī qabūlal-ijābati wa kasyafa lī ‘an asrāri abwābil-inābati ‘alā yadil-‘ulamā'il-mutaqaddimīna wal-a'immati wal-muhtadīna raḍiyallāhu ‘anhum. Żakarū żālika ‘alā anwā‘im-minal-amtsilati wakhtilāfi aṭwālil-abniyati wa lam yūdi‘ūhu minal-ażkāri anwāran, tsumma annahum żakarut-tadrīja ‘alā auzānil-qisṭāsi wa lam yunabbihū ‘alā kaifiyyatil-iḥtirāzi likhtilālil-ḥawāss.",
    "Wa innanī uqaddimu hāżihir-riyāḍātir-rūḥāniyyatal-latī laisa lahā matsīlun khidmatan ilal-mu'minīna fī kulli makānin, wallāhul-muwaffiqu, wa arjū minallāhi ay-yataqabbala hāżal-kitāba wa mā fīhi biqabūlin ḥasanin, wa innanī arjū minallāhi ay-yataqabbalanī ‘indahū wa yataqabbala hāżal-‘amala minal-‘abdil-faqīri ilal-maulal-jalīl. Wa may-yurīdu ay-yaṣila ilal-murādi falyabda' biqirā'atil-fātiḥati lisayyidinā muḥammadin wa ālil-baiti wal-mu'allifi, wa lillāhil-ḥamdu fil-ūlā wal-ākhirah. Al-Mu'allif: Muḥammad Sāmī Muḥammad Aḥmad.",
]

TERJ4 = [
    "Maka Dia (Allah) melimpahkan kepadaku dari cahaya rahmat-Nya penerimaan doa, dan menyingkapkan bagiku rahasia pintu-pintu inabah (kembali kepada Allah) melalui perantara (tangan) para ulama terdahulu dan para imam serta orang-orang yang mendapat petunjuk—semoga Allah meridhai mereka. Mereka telah menyebutkan hal tersebut melalui berbagai macam perumpamaan dan perbedaan tingkatan bangunan (spiritual), namun mereka tidak menyematkan cahaya pada dzikir-dzikir tersebut. Kemudian, mereka menyebutkan tahapan-tahapan atas timbangan keadilan (qistas), namun mereka belum memberikan peringatan tentang cara menjaga diri dari rusaknya panca indera.",
    "Dan sesungguhnya aku mempersembahkan riyadhah-riyadhah ruhaniyah ini—yang tiada tandingannya—sebagai bentuk pelayanan kepada kaum mukminin di setiap tempat. Allah-lah Yang Maha Memberi Taufik. Aku memohon kepada Allah agar menerima kitab ini beserta isinya dengan penerimaan yang baik, dan sungguh aku memohon kepada Allah agar menerimaku di sisi-Nya, serta menerima amal ini dari hamba yang fakir menuju Tuhan Yang Maha Agung. Barangsiapa yang ingin mencapai tujuannya (melalui kitab ini), maka hendaklah ia memulai dengan membaca Al-Fatihah untuk junjungan kita Nabi Muhammad, Ahlul Bait, dan untuk Penulis. Segala puji bagi Allah di dunia dan akhirat.",
    "Penulis: Muhammad Sami Muhammad Ahmad",
]

AR5 = [
    "الْبَابُ الْأَوَّلُ",
    "الرِّيَاضَاتُ الْكُبْرَى",
    "اعْلَمْ أَيُّهَا الطَّالِبُ أَنَّ الْقِيَامَ بِالشَّرَائِعِ النَّبَوِيَّةِ وَمَجْمُوعَةٍ مِنْ حِكْمَةِ اللَّطَائِفِ الرَّبَّانِيَّةِ هُوَ جُودُهُ بَيْنَ الْخَلَائِقِ وَأَنَّ الْعَالَمَ الْفَانِي لَيْسَتْ لَهَا بِهِ عَلَاقَةٌ وَلَا لِصِنْفِ الْتِفَاتٍ لِيَطْمَئِنَّ قَلْبُهُ بِيَوْمِ الطَّلَبِ وَلْيَسْكُنْ جَأْشُ هِمَّتِهِ مِنْ حَيْرَةِ الدَّرْبِ لَكِنْ تَكُونُ الْأَمَارَةُ أَمْرَهُ .",
    "وَمُوَاظَبَةٌ بِإِرَادَتِهَا غَائِبَةً وَقَاهِرَةً وَذَلِكَ لِعِلَّتَيْنِ أَحَدُهُمَا أَنْ يَكُونَ عَقْلُهُ مَحْجُوباً بِالشَّهَوَاتِ وَالْأُخْرَى أَنْ يَكُونَ إِيمَانُهُ مَحْجُوباً تَحْتَ الْمُخَالَفَاتِ فَإِذَا عَمِيَتْ بَصِيرَةُ الْعَقْلِ عَنْ مُشَاهَدَةِ أَرْوَاحِ الْإِيمَانِ بَقِيَتِ النَّفْسُ مَحْجُوبَةً عَلَى النُّورَيْنِ خَالِيَةً مِنَ السِّرَّيْنِ فَيَدْخُلُ السَّاحَةَ الرِّيَاضَةِ بِغَايَةِ الْجِدِّ وَنِهَايَةِ الْجُهْدِ فَيَكُونُ عَلَى ثَلَاثَةِ أَقْسَامٍ .",
    "الْقِسْمُ الْأَوَّلُ :- إِمَّا أَنْ يَقْضِيَ بِهِ عَمَلُهُ لِاخْتِلَالِ الْحَوَاسِّ فَيُخِلَّ بِالْأَعْمَالِ الشَّرْعِيَّةِ وَالْوَظَائِفِ الرَّبَّانِيَّةِ وَرُبَّمَا رَأَى ذَلِكَ أَفْضَلَ وَهَذَا لَا يُفْتَحُ لَهُ فِي سَائِرِ الْأَعْمَالِ بَعْدُ .",
]

LAT5 = [
    "Al-Bābul-Awwal: Ar-Riyāḍātul-Kubrā",
    "I‘lam ayyuhaṭ-ṭālibu annal-qiyāma bisy-syarā'i‘in-nabawiyyati wa majmū‘atim-min ḥikmatil-laṭā'ifir-rabbāniyyati huwa jūduhū bainal-khalā'iqi wa annal-‘ālamal-fānī laisat lahā bihī ‘alāqatun wa lā liṣinfil-tifātin liyaṭma'inna qalbuhū biyaumiṭ-ṭalabi walyaskun ja'syu himmatihī min ḥairatid-darbi lakin takūnul-amāratu amrahū.",
    "Wa muwāẓabatun bi-irādatihā ghā'ibatan wa qāhiratan wa żālika li-'illataini aḥaduhumā ay-yakūna ‘aqluhū maḥjūbam bisy-syahawāti wal-ukhrā ay-yakūna īmānuhū maḥjūban taḥtal-mukhālafāti. Fa-iżā ‘amiyat baṣīratul-‘aqli ‘am-musyāhadati arwāḥil-īmāni baqiyatin-nafsu maḥjūbatan ‘alan-nūraini khāliyatam-minas-sirraini fayadkhulus-sāḥatar-riyāḍati bighāyatil-jiddi wa nihāyatil-juhdi fayakūnu ‘alā tsalātsati aqsām.",
    "Al-Qismul-awwalu: Immā ay-yaqḍiya bihī ‘amaluhū likhtilālil-ḥawāssi fayukhilla bil-a‘mālisy-syar‘iyyati wal-waẓā'ifir-rabbāniyyati wa rubbamā ra'ā żālika afḍala wa hāżā lā yuftaḥu lahū fī sā'iril-a‘māli ba‘du.",
]

TERJ5 = [
    "Bab Pertama: Riyadhah-Riyadhah Besar (Tingkat Tinggi)",
    "Ketahuilah wahai penuntut (ilmu rohani), sesungguhnya mendirikan syariat-syariat kenabian dan himpunan hikmah kelembutan ketuhanan adalah wujud kedermawanan-Nya di antara para makhluk. Dan sesungguhnya alam yang fana ini tidak memiliki ikatan dengannya, dan tidak ada satupun jenis kecondongan (untuk berpaling) agar hatinya tenang pada hari pencarian, dan agar gejolak tekadnya tenang dari kebingungan jalan, akan tetapi tanda-tandanya adalah urusan-Nya (keputusan-Nya).",
    "Dan ketekunan dengan kehendaknya yang gaib lagi memaksa. Hal tersebut disebabkan oleh dua kelemahan (penyakit hati): pertama, akalnya tertutup oleh syahwat; dan kedua, imannya tertutup di bawah berbagai pelanggaran (dosa). Maka apabila mata batin akal telah buta dari menyaksikan ruh-ruh keimanan, tetaplah jiwa itu terhijab dari dua cahaya dan kosong dari dua rahasia. Lalu ia memasuki arena riyadhah dengan puncak kesungguhan dan batas akhir upaya, maka ia akan terbagi menjadi tiga golongan.",
    "Golongan Pertama: Adakalanya amalnya menjerumuskannya pada kerusakan panca indera, sehingga ia merusak amal-amal syariat dan tugas-tugas ketuhanan (rabbaniyyah), dan terkadang ia malah memandang rusaknya indera tersebut sebagai hal yang lebih utama. Orang (golongan) seperti ini tidak akan dibukakan (pintu futuh) baginya pada seluruh amalnya setelah itu.",
]


def has_drawing_el(el) -> bool:
    return bool(el.xpath('.//*[local-name()="drawing"]'))


def make_para(template_el, text: str):
    el = deepcopy(template_el)
    runs = el.findall(qn("w:r"))
    donor = deepcopy(runs[0]) if runs else None
    for child in list(el):
        if child.tag != qn("w:pPr"):
            el.remove(child)
    if donor is not None:
        texts = [t for t in donor.iter() if t.tag == qn("w:t")]
        for t in texts:
            t.text = ""
        if texts:
            texts[0].text = text
            texts[0].set(qn("xml:space"), "preserve")
            for extra in texts[1:]:
                extra.getparent().remove(extra)
        el.append(donor)
    return el


def media_hashes(path: Path) -> list[str]:
    with zipfile.ZipFile(path) as z:
        names = [n for n in z.namelist() if n.startswith("word/media/")]
        return [hashlib.sha256(z.read(n)).hexdigest() for n in names]


def main() -> int:
    doc = Document(TARGET)
    paras = list(doc.paragraphs)
    texts = [p.text for p in paras]
    heads = {}
    for i, t in enumerate(texts):
        m = re.match(r"^=== Halaman (\d+) — (.+) ===$", t.strip())
        if m:
            heads[int(m.group(1))] = i
    h4, h5, h6 = heads[4], heads[5], heads[6]
    els = [p._element for p in paras]

    if any(has_drawing_el(els[i]) for i in range(h4, h6)):
        raise AssertionError("units 4-5 unexpectedly carry an image")

    media_i = next(i for i in range(h4, h5) if texts[i].strip() == "Media:")
    media_el = els[media_i]
    blank5_el = els[h6 - 1]
    if texts[h6 - 1].strip() != "":
        raise AssertionError("expected trailing blank at end of unit 5")

    arab_tpl = deepcopy(els[h4 + 1])
    latin_tpl = deepcopy(els[h4 + 5])
    terj_tpl = deepcopy(els[h4 + 9])
    if not texts[h4 + 5].startswith("Fa afadha"):
        raise AssertionError("latin template mismatch")

    body = els[h4].getparent()
    removed = 0
    for i in list(range(h4 + 1, media_i)) + list(range(h5 + 1, h6 - 1)):
        body.remove(els[i])
        removed += 1

    new4 = [make_para(arab_tpl, t) for t in AR4]
    new4 += [make_para(latin_tpl, t) for t in LAT4]
    new4 += [make_para(terj_tpl, t) for t in TERJ4]
    for el in new4:
        media_el.addprevious(el)

    new5 = [make_para(arab_tpl, t) for t in AR5]
    new5 += [make_para(latin_tpl, t) for t in LAT5]
    new5 += [make_para(terj_tpl, t) for t in TERJ5]
    for el in new5:
        blank5_el.addprevious(el)

    doc.save(TARGET)

    # QA
    out = Document(TARGET)
    ot = [p.text for p in out.paragraphs]
    oh = {}
    for i, t in enumerate(ot):
        m = re.match(r"^=== Halaman (\d+) — (.+) ===$", t.strip())
        if m:
            oh[int(m.group(1))] = i
    if sorted(oh) != list(range(1, 123)):
        raise AssertionError("headings broken")
    u4 = ot[oh[4] + 1 : oh[5]]
    exp4 = AR4 + LAT4 + TERJ4 + [
        "Media:",
        "- Bacaan al-Fatihah",
        "Langkah ke-1:",
        "Membaca al-Fatihah yang dihadiahkan kepada Nabi Muhammad, keluarga beliau, dan pengarang sebelum memulai.",
        "Tujuan:",
        "Menghadiahkan bacaan al-Fatihah kepada Nabi Muhammad, keluarga beliau, dan pengarang sebelum memulai.",
        "",
    ]
    if u4 != exp4:
        raise AssertionError(f"unit4 mismatch: {u4!r}")
    u5 = ot[oh[5] + 1 : oh[6]]
    if u5 != AR5 + LAT5 + TERJ5 + [""]:
        raise AssertionError("unit5 mismatch")
    if len(out.inline_shapes) != 26:
        raise AssertionError("image count changed")
    if media_hashes(TARGET) != media_hashes(SOURCE_FOR_MEDIA):
        raise AssertionError("media hashes/order changed")
    joined = "\n".join(ot)
    for frag in ("فَلْيَبْدَأْ بِقِرَاءَةِ الْفَاتِحَةِ", "بَيْنَ الْخَلَائِقِ",
                 "falyabda' biqirā'atil-fātiḥati", "Penulis: Muhammad Sami Muhammad Ahmad"):
        if frag not in joined:
            raise AssertionError(f"missing fragment: {frag}")

    print(f"removed_old_paras={removed}")
    print(f"total_paragraphs={len(ot)}")
    print(f"bytes={TARGET.stat().st_size}")
    print(f"sha256={hashlib.sha256(TARGET.read_bytes()).hexdigest()}")
    print("H4_5_REVISION_QA_PASS")
    return 0


if __name__ == "__main__":
    sys.exit(main())
