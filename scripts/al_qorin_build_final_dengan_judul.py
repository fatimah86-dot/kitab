#!/usr/bin/env python3
"""Build the final Al-Qorin DOCX *with* restored page headings.

Derived from the audited 2-layer DOCX.  Everything from the label-free
edition stays removed (document title, SAMPUL heading, Layer headings,
``1. Arab:``/``2. Latin:``/``3. Terjemahan:`` labels, footnote and
unclear-print paragraphs), except the 122 ``Halaman X`` Heading-1 lines,
which are kept and reformatted as ``=== Halaman X — ... ===`` in Normal
bold text.  The word ``PARAGRAF`` is stripped from any restored heading.
Arabic, Latin, Terjemahan, Media, Langkah, Tujuan, images, fonts, and
paragraph direction are otherwise left untouched.
"""

from __future__ import annotations

import argparse
import hashlib
import re
import sys
import zipfile
from collections import Counter
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_INPUT = ROOT / "al_qorin/AL_QORIN_123_BERSIH_2_LAYER.docx"
DEFAULT_OUTPUT = ROOT / "al_qorin/AL_QORIN_123_BERSIH_FINAL_DENGAN_JUDUL.docx"

DELETE_STYLES = {"Title", "Heading 2", "Heading 3"}
EXACT_DELETE = {
    "SAMPUL",
    "LAYER 1 - KETERANGAN HALAMAN:",
    "LAYER 1",
    "LAYER 2",
    "PARAGRAF",
    "AL-QORIN WAR RIYADHAH",
    "1. Arab:",
    "2. Latin:",
    "3. Terjemahan:",
}
CONTAINS_DELETE = (
    "tatacara praktik",
    "footnote",
    "[footnote",
    "cetakan/bacaannya tidak jelas",
    "«مصامات»",
    "«المتربرين»",
    "«اخلائق»",
    "«تنبتت»",
    "«الطبايح»",
    "«واعم»",
)

HEADING_RE = re.compile(r"^=== Halaman (\d+) — (.+) ===$")


def normalized(text: str) -> str:
    return " ".join(text.split())


def has_drawing(paragraph) -> bool:
    return bool(paragraph._p.xpath('.//*[local-name()="drawing"]'))


def deletion_reason(paragraph) -> str | None:
    # Never remove a paragraph carrying one of the precision crops.
    if has_drawing(paragraph):
        return None

    if paragraph.style.name in DELETE_STYLES:
        return f"style:{paragraph.style.name}"

    text = normalized(paragraph.text)
    if text in EXACT_DELETE:
        return "exact"

    lowered = paragraph.text.lower()
    for needle in CONTAINS_DELETE:
        if needle.lower() in lowered:
            return f"contains:{needle}"
    return None


def clean_heading_text(text: str) -> str:
    """Remove the banned word PARAGRAF from a restored heading."""
    text = re.sub(r"\bPARAGRAF\b\s*", "", text)
    text = re.sub(r"\s+", " ", text).strip()
    text = text.replace("— —", "—").strip(" —") if text.endswith("—") else text
    return text


def restore_heading(doc, paragraph) -> None:
    """Turn a Heading-1 'Halaman X — ...' line into === ... === Normal bold."""
    new_text = f"=== {clean_heading_text(paragraph.text)} ==="
    paragraph.style = doc.styles["Normal"]
    runs = paragraph.runs
    if runs:
        runs[0].text = new_text
        runs[0].bold = True
        for extra in runs[1:]:
            extra._element.getparent().remove(extra._element)
    else:
        run = paragraph.add_run(new_text)
        run.bold = True


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def media_hashes(path: Path) -> list[str]:
    with zipfile.ZipFile(path) as archive:
        names = [n for n in archive.namelist() if n.startswith("word/media/")]
        return [hashlib.sha256(archive.read(n)).hexdigest() for n in names]


def font_part_count(path: Path) -> int:
    with zipfile.ZipFile(path) as archive:
        return len([n for n in archive.namelist() if n.startswith("word/fonts/")])


def build(input_path: Path, output_path: Path) -> Counter:
    input_path = input_path.resolve()
    output_path = output_path.resolve()
    if input_path == output_path:
        raise ValueError("output must not overwrite its input")
    if output_path.name != "AL_QORIN_123_BERSIH_FINAL_DENGAN_JUDUL.docx":
        raise ValueError("output filename must be AL_QORIN_123_BERSIH_FINAL_DENGAN_JUDUL.docx")

    doc = Document(input_path)
    stats: Counter = Counter()

    # Iterate over a snapshot because paragraph XML elements are removed.
    for paragraph in list(doc.paragraphs):
        if (
            paragraph.style.name == "Heading 1"
            and normalized(paragraph.text) != "SAMPUL"
            and not has_drawing(paragraph)
        ):
            restore_heading(doc, paragraph)
            stats["headings_restored"] += 1
            continue
        reason = deletion_reason(paragraph)
        if reason is not None:
            stats[f"removed:{reason}"] += 1
            element = paragraph._element
            element.getparent().remove(element)

    # Keep package metadata free of the old display title.
    doc.core_properties.title = ""

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)

    # Acceptance checks.
    out = Document(output_path)
    styles = Counter(p.style.name for p in out.paragraphs)
    if any(style in styles for style in DELETE_STYLES | {"Heading 1"}):
        raise AssertionError(f"structural styles remain: {styles}")

    heading_numbers: list[int] = []
    for paragraph in out.paragraphs:
        text = normalized(paragraph.text)
        if text in EXACT_DELETE:
            raise AssertionError(f"exact banned line remains: {text!r}")
        lowered = paragraph.text.lower()
        for needle in CONTAINS_DELETE:
            if needle.lower() in lowered:
                raise AssertionError(f"banned text remains: {needle!r}")
        if "layer 1" in lowered or "layer 2" in lowered:
            raise AssertionError("a Layer label remains")
        if "PARAGRAF" in paragraph.text:
            raise AssertionError("PARAGRAF remains")
        m = HEADING_RE.match(text)
        if m:
            heading_numbers.append(int(m.group(1)))

    if heading_numbers != list(range(1, 123)):
        raise AssertionError(
            f"expected Halaman headings 1..122 in order, got {len(heading_numbers)} headings"
        )

    if len(out.inline_shapes) != 26:
        raise AssertionError(f"expected 26 inline images, got {len(out.inline_shapes)}")
    if media_hashes(output_path) != media_hashes(input_path):
        raise AssertionError("embedded media hashes/order changed")
    if font_part_count(output_path) != 4:
        raise AssertionError("embedded font parts were not preserved")

    media_count = sum(normalized(p.text) == "Media:" for p in out.paragraphs)
    step_count = sum(normalized(p.text).startswith("Langkah ke-") for p in out.paragraphs)
    tujuan_count = sum(normalized(p.text).startswith("Tujuan") for p in out.paragraphs)
    if media_count != 72 or tujuan_count != 72 or step_count < 270:
        raise AssertionError(
            f"practice labels changed unexpectedly: media={media_count} "
            f"steps={step_count} tujuan={tujuan_count}"
        )

    stats["output_paragraphs"] = len(out.paragraphs)
    stats["media_labels"] = media_count
    stats["step_labels"] = step_count
    stats["tujuan_labels"] = tujuan_count
    return stats


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args(argv)

    stats = build(args.input, args.output)
    print(f"input={args.input.resolve()}")
    print(f"output={args.output.resolve()}")
    print(f"stats={dict(stats)}")
    print(f"bytes={args.output.resolve().stat().st_size}")
    print(f"sha256={sha256(args.output.resolve())}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
