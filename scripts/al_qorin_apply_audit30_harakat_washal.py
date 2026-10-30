#!/usr/bin/env python3
"""Lengkapi harakat dan Latin Washal Taqdiriyyah pada 30 unit audit.

Target hanya AL_QORIN_123_FINAL_KELUHAN_SELESAI.docx. Skrip ini tidak
mengubah rasm Arab: huruf dasar setiap token Arab dipertahankan persis;
yang ditambahkan/diseragamkan hanya tanda harakat. Pilihan harakat diambil
berturut-turut dari (1) keselarasan urutan dengan Arab kanonik berharakat
pada unit yang sama, (2) kamus bentuk berharakat dari seluruh korpus
kanonik/FINAL dengan konteks kata tetangga, dan (3) kamus manual/taqdiri
di bawah untuk bentuk yang tidak mempunyai donor yang sepadan.

Latin dibangun ulang dari Arab yang telah berharakat sebagai panduan baca
washal: lam Syamsiyyah diasimilasi, lam Qamariyyah dipertahankan, hamzah
washal artikel gugur sesudah vokal/awalan, tanwin dan i'rab dipertahankan
di tengah rangkaian, serta tasydid digandakan. Paragraf wafaq, gambar,
terjemahan, dan blok praktik tidak diubah. Paragraf Latin khusus yang
memang berupa penanda [ada wafaq] atau baris glyph Arab hitungan pada
Halaman 74 dipertahankan apa adanya.
"""

from __future__ import annotations

import argparse
import hashlib
import importlib.util
import json
import re
import zipfile
from collections import Counter, defaultdict
from difflib import SequenceMatcher
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
TARGET = ROOT / "al_qorin/AL_QORIN_123_FINAL_KELUHAN_SELESAI.docx"
CANON = ROOT / "al_qorin/AL_QORIN_123_BERSIH_HALAMAN.md"
FOTO_JSON = ROOT / "scripts/al_qorin_foto_layer1.json"

TARGET_PAGES = (
    list(range(6, 21))
    + [61, 62, 63, 72, 73, 74]
    + list(range(76, 84))
    + [120]
)

MARKS = set(chr(x) for x in range(0x064B, 0x0660)) | {"\u0670", "\u0653", "\u0654", "\u0655"}
MARKS |= {chr(x) for x in range(0x06D6, 0x06EE)}
TATWEEL = "\u0640"
FATHA, DAMMA, KASRA, SUKUN, SHADDA = "\u064e", "\u064f", "\u0650", "\u0652", "\u0651"
FATHATAN, DAMMATAN, KASRATAN = "\u064b", "\u064c", "\u064d"
SUPER_ALIF = "\u0670"
MADDAH = "\u0653"

NORM_MAP = str.maketrans(
    {
        "أ": "ا",
        "إ": "ا",
        "آ": "ا",
        "ٱ": "ا",
        "ٵ": "ا",
        "ؤ": "و",
        "ئ": "ي",
        "ٸ": "ي",
        "ى": "ي",
        "ة": "ه",
        "ۀ": "ه",
        "ۃ": "ه",
    }
)

# Donor manual ditulis dalam rasm kanonik; injektor hanya memindahkan
# harakatnya ke rasm token FINAL yang sudah ada. Kunci dinormalkan.
MANUAL_RAW = {
    # Partikel, isim isyarah, dan bentuk sangat lazim.
    "ان": "أَنْ",
    "إن": "إِنْ",
    "أن": "أَنَّ",
    "اذا": "إِذَا",
    "إذا": "إِذَا",
    "اذ": "إِذْ",
    "الى": "إِلَى",
    "إلى": "إِلَى",
    "على": "عَلَى",
    "علي": "عَلِيَّ",
    "عن": "عَنْ",
    "من": "مِنْ",
    "ما": "مَا",
    "لا": "لَا",
    "لم": "لَمْ",
    "لن": "لَنْ",
    "قد": "قَدْ",
    "هو": "هُوَ",
    "هى": "هِيَ",
    "هي": "هِيَ",
    "هذا": "هَذَا",
    "هذه": "هَذِهِ",
    "ذلك": "ذَلِكَ",
    "الذى": "الَّذِي",
    "الذي": "الَّذِي",
    "التى": "الَّتِي",
    "التي": "الَّتِي",
    "كل": "كُلِّ",
    "بعد": "بَعْدَ",
    "قبل": "قَبْلَ",
    "عند": "عِنْدَ",
    "غير": "غَيْرَ",
    "مثل": "مِثْلَ",
    "ايها": "أَيُّهَا",
    "يا": "يَا",
    "ثم": "ثُمَّ",
    "او": "أَوْ",
    "أو": "أَوْ",
    "بل": "بَلْ",
    "في": "فِي",
    "فى": "فِي",
    "به": "بِهِ",
    "له": "لَهُ",
    "منه": "مِنْهُ",
    "عليه": "عَلَيْهِ",
    "اليه": "إِلَيْهِ",
    "فيه": "فِيهِ",
    "فيها": "فِيهَا",
    "بها": "بِهَا",
    "لها": "لَهَا",
    "مع": "مَعَ",
    "كيف": "كَيْفَ",
    "حيث": "حَيْثُ",
    "حتى": "حَتَّى",
    "لدى": "لَدَى",
    "لدي": "لَدَيَّ",
    "لديه": "لَدَيْهِ",
    "لديك": "لَدَيْكَ",
    "لديكم": "لَدَيْكُمْ",
    # Bentuk yang sering muncul pada halaman awal dengan rasm cetak ringkas.
    "الله": "اللَّهِ",
    "تعالى": "تَعَالَى",
    "سبحانه": "سُبْحَانَهُ",
    "وسلم": "وَسَلَّمَ",
    "رضى": "رَضِيَ",
    "عنهم": "عَنْهُمْ",
    "عنه": "عَنْهُ",
    "اربعين": "أَرْبَعِينَ",
    "الاربعين": "الْأَرْبَعِينَ",
    "الرياضه": "الرِّيَاضَةِ",
    "الرياضة": "الرِّيَاضَةِ",
    "السالكين": "السَّالِكِينَ",
    "المريدين": "الْمُرِيدِينَ",
    "العارفين": "الْعَارِفِينَ",
    "الاولى": "الْأُولَى",
    "الثانى": "الثَّانِي",
    "الثاني": "الثَّانِي",
    "الثانيه": "الثَّانِيَةِ",
    "الثانية": "الثَّانِيَةِ",
    "الثالث": "الثَّالِثُ",
    "الاول": "الْأَوَّلُ",
    "الاولا": "الْأُولَى",
    "الملائكة": "الْمَلَائِكَةِ",
    "السموات": "السَّمَاوَاتِ",
    "والارض": "وَالْأَرْضِ",
    "الاسماء": "الْأَسْمَاءِ",
    "الحقيقه": "الْحَقِيقَةِ",
    "الحقيقة": "الْحَقِيقَةِ",
    "الحضره": "الْحَضْرَةِ",
    "الحضرة": "الْحَضْرَةِ",
    "الالهيه": "الْإِلَهِيَّةِ",
    "الإلهية": "الْإِلَهِيَّةِ",
    "مرة": "مَرَّةً",
    "مرات": "مَرَّاتٍ",
    # Halaman praktik dan asma'.
    "يكتب": "يُكْتَبُ",
    "تكتب": "تَكْتُبُ",
    "اكتب": "اكْتُبْ",
    "تقرأ": "تَقْرَأُ",
    "تقول": "تَقُولُ",
    "فائدة": "فَائِدَةٌ",
    "باب": "بَابُ",
    "اسكن": "اسْكُنْ",
    "ايها": "أَيُّهَا",
    "كذا": "كَذَا",
    "فلان": "فُلَانٍ",
    "فلانه": "فُلَانَةَ",
    "فلانة": "فُلَانَةَ",
    "عدد": "عَدَدَ",
    "مرتين": "مَرَّتَيْنِ",
    # Sisa audit yang tidak mempunyai donor seunit/global yang sepadan.
    # Bacaan mengikuti konteks dan Latin yang sudah ada; untuk asma'
    # a'jamiyyah dipakai Washal Taqdiriyyah agar dapat disambung.
    "ومافيه": "وَمَافِيهِ",
    "بالحرف": "بِالْحَرْفِ",
    "ولا": "وَلَا",
    "ومثل": "وَمِثْلِ",
    "تفقد": "تَفْقِدُ",
    "مابينه": "مَابَيَّنَهُ",
    "خرطينة": "خَرْطِينَةَ",
    "لحكم": "لِحُكْمِ",
    "فبقنب": "فَبَقْنَبَ",
    "تعاى": "تَعَاى",
    "يايتها": "يَاأَيَّتُهَا",
    "ولاحول": "وَلَاحَوْلَ",
    "وامن": "وَآمَنَ",
    "يس": "يٰسٓ",
    "اتم": "أَتَمَّ",
    "فتحا": "فَتْحًا",
    "استيلا": "اسْتِيلَا",
    "لاعلما": "لَاعِلْمًا",
    "ياذا": "يَاذَا",
    "لايحاط": "لَايُحَاطُ",
    "فلينذكرن": "فَلْيَنْذَكِرَنَّ",
    "مايرد": "مَايَرِدُ",
    "متضجعا": "مُتَضَجِّعًا",
    "لاهم": "لِأَهَمِّ",
    "أثارالملكوت": "أٰثَارَالْمَلَكُوتِ",
    "ادار": "أَدَارَ",
    "يشكل": "يُشْكِلُ",
    "عزوجل": "عَزَّوَجَلَّ",
    "وحولها": "وَحَوْلَهَا",
    "الم": "الٓمٓ",
    "كهيعص": "كٓهيعٓصٓ",
    "بحملته": "بِحَمْلَتِهِ",
    "يرا": "يَرَا",
    "وصوره": "وَصُورَهُ",
    "يكن": "يَكُنْ",
    "كلما": "كُلَّمَا",
    "بامرالله": "بِأَمْرِاللَّهِ",
    "انقلبكم": "انْقَلَبْكُمْ",
    "يادم": "يَادَمُ",
    "طه": "طٰهٓ",
    "لاحول": "لَاحَوْلَ",
    "جبال": "جِبَالٍ",
    "باصرافه": "بِإِصْرَافِهِ",
    "تتأخر": "تَتَأَخَّرْ",
    "تصرفه": "تَصْرِفُهُ",
    "سرعة": "سُرْعَةٍ",
    "تأخرت": "تَأَخَّرْتَ",
    "يضربك": "يَضْرِبْكَ",
    "بالمجمرة": "بِالْمَجْمَرَةِ",
    "فينزل": "فَيَنْزِلُ",
    "كسفا": "كِسْفًا",
    "ولاتامين": "وَلَاتَأْمِينٍ",
    "باسي": "بَاسِي",
    "ياسمم": "يَاسَمَمُ",
    "لاتسرح": "لَاتَسْرَحْ",
    "امين": "أَمِينٌ",
    "ماتكتب": "مَاتَكْتُبُ",
    "وتتعزم": "وَتَتَعَزَّمُ",
    "بطو": "بَطُو",
    "شليش": "شَلْيِيشٍ",
    "الراس": "الرَّاسِ",
    "يه": "يَهْ",
}


def strip_marks(text: str) -> str:
    return "".join(ch for ch in text if ch not in MARKS and ch != TATWEEL)


def norm_skeleton(text: str) -> str:
    return strip_marks(text).translate(NORM_MAP)


def is_arabic_letter(ch: str) -> bool:
    return (
        ("\u0621" <= ch <= "\u064a")
        or ("\u066e" <= ch <= "\u06d3")
        or ch in {"\u0671", "\u06fa", "\u06fc", "\u06ff"}
    )


def is_core_char(ch: str) -> bool:
    return (
        is_arabic_letter(ch)
        or ch in MARKS
        or ch == TATWEEL
        or ch.isdigit()
        or ("\u0660" <= ch <= "\u0669")
        or ("\u06f0" <= ch <= "\u06f9")
    )


def split_token(token: str) -> tuple[str, str, str]:
    idx = [i for i, ch in enumerate(token) if is_core_char(ch)]
    if not idx:
        return token, "", ""
    a, b = idx[0], idx[-1] + 1
    return token[:a], token[a:b], token[b:]


def segments(core: str) -> list[dict]:
    out: list[dict] = []
    pending_marks: list[str] = []
    for ch in core:
        if ch in MARKS:
            if out:
                out[-1]["marks"].append(ch)
            else:
                pending_marks.append(ch)
        elif ch == TATWEEL:
            if out:
                out[-1]["tatweel"] = out[-1].get("tatweel", "") + ch
        else:
            out.append({"base": ch, "marks": pending_marks.copy(), "tatweel": ""})
            pending_marks.clear()
    if pending_marks and out:
        out[-1]["marks"].extend(pending_marks)
    return out


def render_segments(segs: list[dict]) -> str:
    return "".join(s["base"] + "".join(s["marks"]) + s.get("tatweel", "") for s in segs)


def mark_count(text: str) -> int:
    return sum(ch in MARKS for ch in text)


def letter_count(text: str) -> int:
    return sum(is_arabic_letter(ch) for ch in text)


def token_quality(core: str) -> float:
    lc = letter_count(core)
    return mark_count(core) / lc if lc else 0.0


def inject(current_core: str, donor_core: str) -> tuple[str, float]:
    cur = segments(current_core)
    don = segments(donor_core)
    if not cur or not don:
        return current_core, 0.0
    cur_norm = [s["base"].translate(NORM_MAP) for s in cur]
    don_norm = [s["base"].translate(NORM_MAP) for s in don]
    assigned: dict[int, list[str]] = {}

    def put(ci: int, di: int) -> None:
        marks = list(don[di]["marks"])
        # Donor آ terhadap ا pada rasm FINAL dinyatakan dengan maddah,
        # tanpa mengganti huruf dasarnya.
        if don[di]["base"] == "آ" and cur[ci]["base"] == "ا" and MADDAH not in marks:
            marks.append(MADDAH)
        assigned[ci] = marks

    sm = SequenceMatcher(None, cur_norm, don_norm, autojunk=False)
    for tag, i1, i2, j1, j2 in sm.get_opcodes():
        if tag == "equal":
            for k in range(i2 - i1):
                put(i1 + k, j1 + k)
        elif tag == "replace":
            for k in range(min(i2 - i1, j2 - j1)):
                put(i1 + k, j1 + k)
    for i, seg in enumerate(cur):
        if i in assigned:
            # Pertahankan harakat lama hanya bila donor tidak memberi tanda
            # pada konsonan. Kosongnya tanda donor pada huruf mad adalah
            # informasi: tanwin lama di atas alif tanwin, misalnya, harus
            # hilang agar tidak terjadi tanwin ganda.
            if assigned[i]:
                seg["marks"] = assigned[i]
            elif seg["base"] in {"ا", "ٱ", "آ", "ى"}:
                seg["marks"] = []
            elif seg["base"] in {"و", "ي"} and any(
                mark in {FATHATAN, DAMMATAN, KASRATAN} for mark in seg["marks"]
            ):
                seg["marks"] = []
            elif not seg["marks"]:
                seg["marks"] = []
    score = sm.ratio()
    return render_segments(cur), score


MANUAL = {norm_skeleton(k): split_token(v)[1] for k, v in MANUAL_RAW.items()}
# Kunci eksak mempertahankan perbedaan rasm yang hilang saat normalisasi,
# terutama على (alā) berhadapan dengan علي (Aliyy/'alayya).
MANUAL_EXACT = {strip_marks(k): split_token(v)[1] for k, v in MANUAL_RAW.items()}


def fallback_short_glyph(core: str) -> str:
    """Beri tanda baca minimal pada glyph hijaiyah satu/dua huruf.

    Ini hanya untuk deret huruf lepas (misalnya Halaman 120), bukan untuk
    menebak kata biasa. Huruf tunggal dibaca sebagai glyph bersukun;
    alif/hamzah dan wawu penghubung diberi fathah. Dua glyph dibaca
    taqdiri fathah+sukun agar tidak ada huruf tanpa tanda.
    """
    segs = segments(core)
    if not segs or len(segs) > 2 or token_quality(core) >= 0.42:
        return core
    if len(segs) == 1:
        base = segs[0]["base"]
        segs[0]["marks"] = [FATHA] if base in {"ا", "أ", "إ", "آ", "ٱ", "و", "ي"} else [SUKUN]
    else:
        if not segs[0]["marks"]:
            segs[0]["marks"] = [FATHA]
        if not segs[1]["marks"]:
            segs[1]["marks"] = [SUKUN]
    return render_segments(segs)


def paragraph_tokens(text: str) -> list[str]:
    return text.split()


def canonical_arabic() -> dict[int, list[str]]:
    lines = CANON.read_text(encoding="utf-8").splitlines()
    units: dict[int, list[str]] = {}
    page = None
    in_arab = False
    cur: list[str] = []
    buf: list[str] = []

    def flush_buf() -> None:
        nonlocal buf
        if buf and page is not None:
            txt = " ".join(x.strip() for x in buf).strip()
            if txt:
                cur.append(txt)
        buf = []

    def flush_unit() -> None:
        nonlocal cur
        if page is not None and cur:
            units[page] = cur
        cur = []

    for line in lines:
        m = re.match(r"^## (?:SAMPUL|Halaman (\d+) —)", line)
        if m:
            flush_buf()
            flush_unit()
            page = int(m.group(1)) if m.group(1) else 0
            in_arab = False
            continue
        if line.startswith("### 1. Arab"):
            in_arab = True
            continue
        if in_arab and line.startswith("### "):
            flush_buf()
            in_arab = False
            continue
        if in_arab:
            if line.strip():
                buf.append(line)
            else:
                flush_buf()
    flush_buf()
    flush_unit()
    return units


def heading_map(texts: list[str]) -> dict[int, int]:
    out = {}
    for i, text in enumerate(texts):
        m = re.match(r"^=== Halaman (\d+) — (.+) ===$", text.strip())
        if m:
            out[int(m.group(1))] = i
    return out


def load_module(name: str):
    path = ROOT / f"scripts/al_qorin_apply_{name}_user_revision.py"
    spec = importlib.util.spec_from_file_location(name, path)
    mod = importlib.util.module_from_spec(spec)
    assert spec.loader is not None
    spec.loader.exec_module(mod)
    return mod


def target_index_map(doc: Document) -> dict[int, dict[str, list[int]]]:
    texts = [p.text for p in doc.paragraphs]
    heads = heading_map(texts)
    result: dict[int, dict[str, list[int]]] = {}

    foto = json.loads(FOTO_JSON.read_text(encoding="utf-8"))
    for page in range(6, 21):
        entry = foto[str(page)]
        ac = len([x for x in entry["Arab"].split("\n\n") if x.strip()])
        lc = len([x for x in entry["Latin"].split("\n\n") if x.strip()])
        s = heads[page] + 1
        result[page] = {
            "ar": list(range(s, s + ac)),
            "la": list(range(s + ac, s + ac + lc)),
        }

    # Unit tanpa gambar dari modul revisi pengguna: AR lalu LAT berurutan.
    for modname, pages in [
        ("h59_61", [61]),
        ("h62_63", [62, 63]),
        ("h80_81", [80, 81]),
    ]:
        mod = load_module(modname)
        for page in pages:
            ac = len(getattr(mod, f"AR{page}"))
            lc = len(getattr(mod, f"LAT{page}"))
            s = heads[page] + 1
            result[page] = {
                "ar": list(range(s, s + ac)),
                "la": list(range(s + ac, s + ac + lc)),
            }
            assert texts[s : s + ac] == getattr(mod, f"AR{page}"), page
            assert texts[s + ac : s + ac + lc] == getattr(mod, f"LAT{page}"), page

    # Unit bergambar: build_sequence memetakan setiap paragraf, termasuk
    # posisi gambar yang harus tetap.
    for modname, pages in [
        ("h72_73", [72, 73]),
        ("h74_75", [74]),
        ("h76_77", [76, 77]),
        ("h78_79", [78, 79]),
        ("h82_83", [82, 83]),
    ]:
        mod = load_module(modname)
        for page in pages:
            seq = mod.build_sequence(page)
            s = heads[page] + 1
            ar: list[int] = []
            la: list[int] = []
            for off, (kind, expected) in enumerate(seq):
                idx = s + off
                if kind == "img":
                    assert doc.paragraphs[idx]._p.xpath('.//*[local-name()="drawing"]'), (page, idx)
                else:
                    assert texts[idx] == expected, (page, idx, kind)
                    if kind == "ar":
                        ar.append(idx)
                    elif kind == "la":
                        la.append(idx)
            result[page] = {"ar": ar, "la": la}

    # Halaman 120: enam paragraf Arab lalu dua paragraf Latin; gambar dan
    # blok praktik sesudah terjemahan tidak tersentuh.
    s = heads[120] + 1
    result[120] = {"ar": list(range(s, s + 6)), "la": [s + 6, s + 7]}
    assert texts[s].startswith("الْفَصْلُ")
    assert texts[s + 6].startswith("Al-fashluts")
    return result


class Vocalizer:
    def __init__(self, doc: Document, canon: dict[int, list[str]]):
        self.canon = canon
        self.occurrences: dict[str, list[tuple[str, str, str, int]]] = defaultdict(list)
        self.frequency: Counter[tuple[str, str]] = Counter()
        self._harvest_canonical()
        self._harvest_final(doc)

    def _add_sequence(self, texts: list[str], priority: int) -> None:
        cores: list[str] = []
        for text in texts:
            for tok in paragraph_tokens(text):
                _, core, _ = split_token(tok)
                if letter_count(core):
                    cores.append(core)
        keys = [norm_skeleton(c) for c in cores]
        for i, core in enumerate(cores):
            if token_quality(core) < 0.22 and mark_count(core) < 2:
                continue
            prev = keys[i - 1] if i else ""
            nxt = keys[i + 1] if i + 1 < len(keys) else ""
            self.occurrences[keys[i]].append((core, prev, nxt, priority))
            self.frequency[(keys[i], core)] += 1

    def _harvest_canonical(self) -> None:
        for page in sorted(self.canon):
            self._add_sequence(self.canon[page], 2)

    def _harvest_final(self, doc: Document) -> None:
        for para in doc.paragraphs:
            text = para.text
            ar = len(re.findall(r"[\u0600-\u06FF]", text))
            la = len(re.findall(r"[A-Za-z]", text))
            if ar >= 3 and ar >= la:
                self._add_sequence([text], 1)

    def global_donor(self, key: str, prev: str, nxt: str) -> str | None:
        cands = self.occurrences.get(key, [])
        if not cands:
            return None

        def score(item):
            core, p, n, pr = item
            return (
                (8 if p and p == prev else 0)
                + (8 if n and n == nxt else 0)
                + pr
                + min(4, self.frequency[(key, core)])
                + token_quality(core)
            )

        return max(cands, key=score)[0]

    def same_page_map(self, page: int, current_texts: list[str]) -> dict[int, str]:
        current: list[str] = []
        for text in current_texts:
            for tok in paragraph_tokens(text):
                _, core, _ = split_token(tok)
                if letter_count(core):
                    current.append(core)
        donor_texts = self.canon.get(page, [])
        donor: list[str] = []
        for text in donor_texts:
            for tok in paragraph_tokens(text):
                _, core, _ = split_token(tok)
                if letter_count(core):
                    donor.append(core)
        sm = SequenceMatcher(
            None,
            [norm_skeleton(x) for x in current],
            [norm_skeleton(x) for x in donor],
            autojunk=False,
        )
        out: dict[int, str] = {}
        for tag, i1, i2, j1, j2 in sm.get_opcodes():
            if tag == "equal":
                for k in range(i2 - i1):
                    out[i1 + k] = donor[j1 + k]
            elif tag == "replace" and (i2 - i1) == (j2 - j1):
                for k in range(i2 - i1):
                    a, b = norm_skeleton(current[i1 + k]), norm_skeleton(donor[j1 + k])
                    if SequenceMatcher(None, a, b, autojunk=False).ratio() >= 0.48:
                        out[i1 + k] = donor[j1 + k]
        return out

    def vocalize_page(self, page: int, texts: list[str]) -> tuple[list[str], list[tuple[str, str, float]]]:
        same = self.same_page_map(page, texts)
        # Indeks token Arab lintas paragraf pada unit ini.
        flat: list[tuple[int, int, str, str, str]] = []  # pidx, tidx, pre, core, suf
        split_texts: list[list[str]] = []
        for pi, text in enumerate(texts):
            toks = text.split(" ")
            split_texts.append(toks)
            for ti, tok in enumerate(toks):
                pre, core, suf = split_token(tok)
                if letter_count(core):
                    flat.append((pi, ti, pre, core, suf))
        keys = [norm_skeleton(x[3]) for x in flat]
        problems: list[tuple[str, str, float]] = []
        for pos, (pi, ti, pre, core, suf) in enumerate(flat):
            key = keys[pos]
            prev = keys[pos - 1] if pos else ""
            nxt = keys[pos + 1] if pos + 1 < len(keys) else ""
            candidates = []
            if same.get(pos):
                candidates.append(same[pos])
            manual_exact = MANUAL_EXACT.get(strip_marks(core))
            if manual_exact:
                candidates.append(manual_exact)
            global_candidate = self.global_donor(key, prev, nxt)
            if global_candidate:
                candidates.append(global_candidate)
            if MANUAL.get(key) and MANUAL[key] not in candidates:
                candidates.append(MANUAL[key])

            new_core, sim = core, 0.0
            best_score = token_quality(core) * 1.5
            for rank, donor in enumerate(candidates):
                candidate, candidate_sim = inject(core, donor)
                quality = token_quality(candidate)
                # Donor pertama yang rasmnya tepat dan sudah cukup
                # berharakat dipakai langsung. Ini mencegah kamus umum
                # yang kebetulan mempunyai tasydid mengalahkan bacaan
                # kontekstual seunit.
                if candidate_sim >= 0.999 and quality >= 0.42:
                    new_core, sim = candidate, candidate_sim
                    break
                priority_bonus = (0.08, 0.05, 0.02, 0.0)[min(rank, 3)]
                score = quality * 1.5 + candidate_sim + priority_bonus
                if score > best_score:
                    new_core, sim, best_score = candidate, candidate_sim, score
            new_core = fallback_short_glyph(new_core)
            split_texts[pi][ti] = pre + new_core + suf
            if token_quality(new_core) < 0.42:
                problems.append((core, new_core, token_quality(new_core)))
        return [" ".join(x) for x in split_texts], problems


SUN_LETTERS = set("تثدذرزسشصضطظلن")
LATIN_CONS = {
    "ء": "'",
    "أ": "'",
    "إ": "'",
    "ؤ": "'",
    "ئ": "'",
    "ب": "b",
    "ت": "t",
    "ث": "ts",
    "ج": "j",
    "ح": "ḥ",
    "خ": "kh",
    "د": "d",
    "ذ": "ż",
    "ر": "r",
    "ز": "z",
    "س": "s",
    "ش": "sy",
    "ص": "ṣ",
    "ض": "ḍ",
    "ط": "ṭ",
    "ظ": "ẓ",
    "ع": "‘",
    "غ": "gh",
    "ف": "f",
    "ق": "q",
    "ك": "k",
    "ل": "l",
    "م": "m",
    "ن": "n",
    "ه": "h",
    "ة": "t",
    "و": "w",
    "ي": "y",
}
LETTER_TOKEN_READINGS = {
    "الم": "alif lām mīm",
    "الر": "alif lām rā",
    "المر": "alif lām mīm rā",
    "المص": "alif lām mīm ṣād",
    "كهيعص": "kāf hā yā ‘ain ṣād",
    "يس": "yā sīn",
    "طه": "ṭā hā",
    "حم": "ḥā mīm",
    "عسق": "‘ain sīn qāf",
    "طس": "ṭā sīn",
    "طسم": "ṭā sīn mīm",
    "ص": "ṣād",
    "ق": "qāf",
    "ن": "nūn",
}
DIGIT_TRANS = str.maketrans("٠١٢٣٤٥٦٧٨٩۰۱۲۳۴۵۶۷۸۹", "01234567890123456789")
PUNCT_TRANS = str.maketrans({"،": ",", "؛": ";", "؟": "?", "«": '"', "»": '"'})


def vowel_mark(seg: dict) -> str | None:
    for mark in (FATHATAN, DAMMATAN, KASRATAN, FATHA, DAMMA, KASRA):
        if mark in seg["marks"]:
            return mark
    return None


def simple_vowel(seg: dict, default: str = "a") -> str:
    mark = vowel_mark(seg)
    return {FATHA: "a", DAMMA: "u", KASRA: "i"}.get(mark, default)


def double_cons(cons: str) -> str:
    return cons + cons


def allah_reading(core: str, pause: bool) -> str | None:
    key = norm_skeleton(core)
    stems = {
        "الله": "allāh",
        "بالله": "billāh",
        "لله": "lillāh",
        "والله": "wallāh",
        "فالله": "fallāh",
        "تالله": "tallāh",
        "اللهم": "allāhumma",
    }
    if key not in stems:
        return None
    stem = stems[key]
    if key == "اللهم" or pause:
        return stem
    segs = segments(core)
    if segs:
        mark = vowel_mark(segs[-1])
        suffix = {FATHA: "a", DAMMA: "u", KASRA: "i"}.get(mark)
        if suffix:
            return stem + suffix
    return stem


def render_core_latin(
    core: str,
    *,
    pause: bool,
    after_sound: bool,
    preserve_single: bool = False,
) -> tuple[str, bool, str]:
    """Transliterasi satu token; hasil kedua menandai artikel washal."""
    if not letter_count(core):
        return core.translate(DIGIT_TRANS), False, ""
    key = norm_skeleton(core)
    if preserve_single and (
        len(segments(core)) == 1 or norm_skeleton(core) == "صع"
    ):
        return strip_marks(core), False, segments(core)[0]["base"].translate(NORM_MAP)
    special = allah_reading(core, pause)
    if special is not None:
        return special, False, "ل"
    if key == "تعالي":
        return "ta‘ālā", False, "ت"
    if key == "عزوجل":
        return "‘azza wa jalla", False, "ع"
    exact_skeleton = strip_marks(core)
    segs_for_special = segments(core)
    if key == "الي" and len(segs_for_special) == 3 and SHADDA not in segs_for_special[-1]["marks"]:
        return "ilā", False, "ا"
    if key == "الا" and len(segs_for_special) == 3:
        # Rasm yang sama memuat إلّا dan ألا; tanda pada lam/alif membedakan.
        if KASRA in segs_for_special[0]["marks"] or SHADDA in segs_for_special[1]["marks"]:
            return "illā", False, "ا"
        return "alā", False, "ا"
    relative_readings = {
        "الذي": "allażī",
        "التي": "allatī",
        "الذين": "allażīna",
        "اللذين": "allażaini",
        "اللاتي": "allātī",
        "اللايي": "allā'ī",
    }
    if key in relative_readings:
        return relative_readings[key], False, "ل"
    if exact_skeleton == "على" or (
        key == "علي"
        and len(segs_for_special) == 3
        and FATHA in segs_for_special[1]["marks"]
        and SHADDA not in segs_for_special[2]["marks"]
    ):
        return "‘alā", False, "ع"
    ilayhi_readings = {
        "اليه": "ilaihi",
        "اليها": "ilaihā",
        "اليك": "ilaika",
        "اليكم": "ilaikum",
        "الينا": "ilainā",
        "الي": "ilaiya",
        "واليه": "wailaihi",
        "واليها": "wailaihā",
        "واليك": "wailaika",
        "واليكم": "wailaikum",
        "والينا": "wailainā",
    }
    if key in ilayhi_readings and not (
        key == "الي" and SHADDA not in segs_for_special[-1]["marks"]
    ):
        return ilayhi_readings[key], False, "ا"
    if key == "هو":
        return "huwa", False, "ه"
    if key == "هي":
        return "hiya", False, "ه"
    if key in LETTER_TOKEN_READINGS:
        return LETTER_TOKEN_READINGS[key], False, segments(core)[0]["base"].translate(NORM_MAP)

    segs = segments(core)
    out = ""
    offset = 0
    prefix_out = ""
    join_prev = False
    sun_index = -1

    # Awalan satu huruf hanya dipisah secara fonetis sebelum artikel;
    # pada kata biasa urutan hurufnya tetap dibaca lurus.
    if (
        len(segs) >= 3
        and segs[0]["base"] in {"و", "ف", "ب", "ل", "ك"}
        and segs[1]["base"] in {"ا", "ٱ", "آ"}
        and segs[2]["base"] == "ل"
    ):
        p = segs[0]
        default = {"و": "a", "ف": "a", "ب": "i", "ل": "i", "ك": "a"}[p["base"]]
        prefix_out = LATIN_CONS[p["base"]] + simple_vowel(p, default)
        out += prefix_out
        offset = 1

    # Artikel: Syamsiyyah diasimilasi; Qamariyyah mempertahankan lam.
    if (
        len(segs) - offset >= 3
        and segs[offset]["base"] in {"ا", "ٱ", "آ"}
        and segs[offset + 1]["base"] == "ل"
    ):
        sun_index = offset + 2
        next_base = segs[sun_index]["base"].translate(NORM_MAP)
        joined = after_sound or bool(prefix_out)
        if next_base in SUN_LETTERS:
            cons = LATIN_CONS.get(segs[sun_index]["base"], next_base)
            out += (cons if joined else "a" + cons) + "-"
        else:
            out += ("l" if joined else "al") + "-"
        join_prev = joined and not prefix_out
        offset += 2

    first_base = ""
    if offset < len(segs):
        first_base = segs[offset]["base"].translate(NORM_MAP)

    i = offset
    while i < len(segs):
        seg = segs[i]
        base = seg["base"]
        marks = seg["marks"]
        nxt = segs[i + 1] if i + 1 < len(segs) else None

        if base in {"ا", "ٱ", "آ"}:
            # Alif farqah sesudah wawu jama' tidak dibunyikan.
            if i == len(segs) - 1 and i > offset and segs[i - 1]["base"] == "و" and not marks:
                i += 1
                continue
            alif_mark = vowel_mark(seg)
            if base == "آ" or MADDAH in marks or SUPER_ALIF in marks:
                out += "ā"
                i += 1
                continue
            if alif_mark in {FATHATAN, DAMMATAN, KASRATAN}:
                if not (pause and i == len(segs) - 1 and alif_mark != FATHATAN):
                    out += {FATHATAN: "an", DAMMATAN: "un", KASRATAN: "in"}[alif_mark]
                elif alif_mark == FATHATAN:
                    out += "ā"
                i += 1
                continue
            if alif_mark in {FATHA, DAMMA, KASRA}:
                # Ibn/ism adalah hamzah washal: gugur di tengah rangkaian.
                # Hanya bentuk eksak berikut; jangan sampai ابنية/أسماء
                # ikut kehilangan alif awalnya.
                wasl_exact = {
                    "ابن", "ابنه", "ابنها", "ابني", "ابنك", "ابنكم",
                    "اسم", "اسمه", "اسمها", "اسمي", "اسمك", "اسمكم",
                }
                if i == offset and after_sound and key in wasl_exact:
                    i += 1
                    continue
                # Alif berharakat pendek yang diikuti huruf mad adalah
                # satu bunyi panjang/diftong, bukan vokal pendek + mad.
                if nxt is not None and vowel_mark(nxt) is None and SHADDA not in nxt["marks"]:
                    combo = None
                    if alif_mark == FATHA and nxt["base"] in {"ا", "ٱ", "آ", "ى"}:
                        combo = "ā"
                    elif alif_mark == FATHA and nxt["base"] == "و":
                        combo = "au"
                    elif alif_mark == FATHA and nxt["base"] == "ي":
                        combo = "ai"
                    elif alif_mark == DAMMA and nxt["base"] == "و":
                        combo = "ū"
                    elif alif_mark == KASRA and nxt["base"] in {"ي", "ى"}:
                        combo = "ī"
                    if combo is not None:
                        out += combo
                        i += 2
                        continue
                out += simple_vowel(seg)
            elif i == offset:
                wasl_exact = {
                    "ابن", "ابنه", "ابنها", "ابني", "ابنك", "ابنكم",
                    "اسم", "اسمه", "اسمها", "اسمي", "اسمك", "اسمكم",
                }
                if not (after_sound and key in wasl_exact):
                    out += "a"
            else:
                out += "ā"
            i += 1
            continue
        if base == "ى":
            ya_mark = vowel_mark(seg)
            if ya_mark is None and SHADDA not in marks:
                # ى tanpa tanda sesudah fathah adalah alif maqsūrah.
                # (Sesudah kasrah/dammah ia sudah dikonsumsi sebagai
                # mad oleh huruf sebelumnya.)
                out += "ā"
            else:
                # ى yang bertanda/bertasydid berfungsi sebagai konsonan y,
                # misalnya رَضِيَ, السَّنِيِّ, dan يُصَلِّي.
                out += "yy" if SHADDA in marks else "y"
                if ya_mark in {FATHA, DAMMA, KASRA}:
                    if not (pause and i == len(segs) - 1):
                        out += {FATHA: "a", DAMMA: "u", KASRA: "i"}[ya_mark]
                elif ya_mark in {FATHATAN, DAMMATAN, KASRATAN}:
                    if pause and i == len(segs) - 1:
                        out += "ā" if ya_mark == FATHATAN else ""
                    else:
                        out += {FATHATAN: "an", DAMMATAN: "un", KASRATAN: "in"}[ya_mark]
            i += 1
            continue
        if base == "و" and vowel_mark(seg) is None and SHADDA not in marks:
            out += "w" if i == offset else "ū"
            i += 1
            continue
        if base == "ي" and vowel_mark(seg) is None and SHADDA not in marks:
            out += "y" if i == offset else "ī"
            i += 1
            continue

        cons = LATIN_CONS.get(base, "")
        if base in {"أ", "إ"} and i == offset:
            # Gaya baca yang dipakai berkas ini tidak menulis apostrof
            # hamzah pada awal kata atau tepat sesudah artikel (al-asmā').
            cons = ""
        if base == "ة":
            cons = "h" if (pause or vowel_mark(seg) is None) and i == len(segs) - 1 else "t"
        if SHADDA in marks and i != sun_index and cons:
            cons = double_cons(cons)
        out += cons

        if SUPER_ALIF in marks:
            out += "ā"
            i += 1
            continue

        mark = vowel_mark(seg)
        skip_next = False
        is_final = i == len(segs) - 1
        if mark in {FATHATAN, DAMMATAN, KASRATAN}:
            if mark == FATHATAN and nxt is not None and nxt["base"] in {"ا", "ٱ"} and not nxt["marks"]:
                skip_next = True
                is_final = i + 1 == len(segs) - 1
            if pause and is_final:
                out += "ā" if mark == FATHATAN else ""
            else:
                out += {FATHATAN: "an", DAMMATAN: "un", KASRATAN: "in"}[mark]
        elif mark == FATHA:
            if nxt is not None and nxt["base"] in {"ا", "ٱ", "آ", "ى"} and vowel_mark(nxt) is None:
                out += "ā"
                skip_next = True
            elif nxt is not None and nxt["base"] == "و" and vowel_mark(nxt) is None and SHADDA not in nxt["marks"]:
                out += "au"
                skip_next = True
            elif nxt is not None and nxt["base"] == "ي" and vowel_mark(nxt) is None and SHADDA not in nxt["marks"]:
                out += "ai"
                skip_next = True
            elif not (pause and is_final):
                out += "a"
        elif mark == KASRA:
            if nxt is not None and nxt["base"] in {"ي", "ى"} and vowel_mark(nxt) is None and SHADDA not in nxt["marks"]:
                out += "ī"
                skip_next = True
            elif not (pause and is_final):
                out += "i"
        elif mark == DAMMA:
            if nxt is not None and nxt["base"] in {"و", "ى"} and vowel_mark(nxt) is None and SHADDA not in nxt["marks"]:
                out += "ū"
                skip_next = True
            elif not (pause and is_final):
                out += "u"
        i += 2 if skip_next else 1

    out = out.lstrip("'")
    out = out.replace("aā", "ā") if out.startswith("wā") is False else out
    return out, join_prev, first_base


def final_nun_kind(core: str) -> bool:
    segs = segments(core)
    if not segs:
        return False
    last = segs[-1]
    if vowel_mark(last) in {FATHATAN, DAMMATAN, KASRATAN}:
        return True
    if last["base"] in {"ا", "ٱ"} and len(segs) >= 2:
        return vowel_mark(segs[-2]) == FATHATAN
    return last["base"] == "ن" and SUKUN in last["marks"]


def capitalize_sentences(text: str) -> str:
    chars = list(text)
    cap = True
    for i, ch in enumerate(chars):
        if cap and ch.isalpha():
            chars[i] = ch.upper()
            cap = False
        elif ch in ".!?":
            cap = True
        elif not ch.isspace() and ch not in {'"', "'", "‘", "(", "["}:
            cap = False
    text = "".join(chars)
    text = re.sub(r"(?<![A-Za-zĀ-Žā-ž])allāh", "Allāh", text)
    text = re.sub(r"(?<![A-Za-zĀ-Žā-ž])muḥammad", "Muḥammad", text)
    return text


def transliterate_paragraph(text: str, *, preserve_single: bool = False) -> str:
    raw_tokens = [x for x in text.split(" ") if x != ""]
    parsed = [split_token(x) for x in raw_tokens]
    result = ""
    prev_core = ""
    sentence_end = True

    for idx, (raw, (pre, core, suf)) in enumerate(zip(raw_tokens, parsed)):
        if not letter_count(core):
            rendered = raw.translate(DIGIT_TRANS).translate(PUNCT_TRANS)
            if rendered in {".", ",", ":", ";", "!", "?", "..."}:
                result += rendered
                sentence_end = rendered in {".", "!", "?"}
            else:
                if result and not result.endswith((" ", "(", '"')):
                    result += " "
                result += rendered
                # Sesudah tanda pisah, artikel dibaca dari awal (as-,
                # bukan ditempel pada tanda pisah).
                sentence_end = rendered in {"-", "—", "–", ":"}
            prev_core = ""
            continue

        next_is_period = idx + 1 < len(raw_tokens) and raw_tokens[idx + 1] in {".", "...", "!", "؟"}
        pause = bool(re.search(r"[.!?؟]$", suf)) or next_is_period or idx == len(raw_tokens) - 1
        after_sound = bool(result) and not sentence_end
        rendered, join_prev, first_base = render_core_latin(
            core,
            pause=pause,
            after_sound=after_sound,
            preserve_single=preserve_single,
        )
        token_out = pre.translate(PUNCT_TRANS) + rendered + suf.translate(DIGIT_TRANS).translate(PUNCT_TRANS)

        assimilate = {"ي": "y", "ر": "r", "م": "m", "ل": "l", "و": "w", "ن": "n", "ب": "m"}
        do_idgham = (
            bool(result)
            and not sentence_end
            and bool(prev_core)
            and final_nun_kind(prev_core)
            and first_base in assimilate
            and result.endswith("n")
            and not join_prev
        )
        if join_prev and result:
            result += token_out
        elif do_idgham:
            result = result[:-1] + assimilate[first_base] + "-" + token_out
        else:
            if result and not result.endswith((" ", "(", '"')):
                result += " "
            result += token_out
        sentence_end = bool(re.search(r"[.!?؟]$", suf))
        prev_core = core

    result = re.sub(r"\s+([,.;:!?])", r"\1", result)
    result = re.sub(r"\(\s+", "(", result)
    result = re.sub(r"\s+\)", ")", result)
    return capitalize_sentences(result.strip())


LATIN_CLEANUPS = [
    ("wal-ceyna", "wal-‘ayna"),
    ("ratْqan", "ratqan"),
    ("lā تَżarnī", "lā teżarnī"),
    ("ْṭiflan", "ṭiflan"),
    ("al-‘arَāf", "al-‘arāf"),
    ("‘aẓīم", "‘aẓīm"),
    ("aradtal-عمل", "aradtal-‘amala"),
    ("wal-‘ādِيāti", "wal-‘ādiyāti"),
    ("حبsuka", "ḥabsuka"),
    ("ikhwāναν", "ikhwānan"),
    ("wasy-sha‘ri", "wasy-sya‘ri"),
    ("bin-naf-ḫi", "bin-nafkhi"),
]


def clean_existing_latin(text: str) -> str:
    if text.strip() in {"[ada wafaq]", "( [ada wafaq] )"}:
        return text
    for old, new in LATIN_CLEANUPS:
        text = text.replace(old, new)
    # Sisa harakat Arab yang menempel pada kata Latin tidak mempunyai
    # fungsi baca dan merusak pencarian/teks.
    text = "".join(ch for ch in text if ch not in MARKS)
    return text


def set_para_text(paragraph, text: str) -> None:
    text_nodes = paragraph._p.findall(".//" + qn("w:t"))
    if not text_nodes:
        paragraph.add_run(text)
        return
    text_nodes[0].text = text
    text_nodes[0].set(qn("xml:space"), "preserve")
    for node in text_nodes[1:]:
        node.text = ""


def media_hashes(path: Path) -> list[str]:
    with zipfile.ZipFile(path) as z:
        return [
            hashlib.sha256(z.read(name)).hexdigest()
            for name in z.namelist()
            if name.startswith("word/media/")
        ]


def image_indices(doc: Document) -> list[int]:
    return [
        i
        for i, p in enumerate(doc.paragraphs)
        if p._p.xpath('.//*[local-name()="drawing"]')
    ]


def build_revisions(doc: Document, index_map, vocalizer):
    revised_text: dict[int, str] = {}
    revised_arabic: dict[int, list[str]] = {}
    problems_total = 0

    for page in TARGET_PAGES:
        ar_idx = index_map[page]["ar"]
        old_ar = [doc.paragraphs[i].text for i in ar_idx]
        new_ar, problems = vocalizer.vocalize_page(page, old_ar)
        problems_total += len(problems)
        assert not problems, (page, problems[:10])
        for old, new in zip(old_ar, new_ar):
            assert strip_marks(old) == strip_marks(new), (page, old, new)
        revised_arabic[page] = new_ar
        for idx, text in zip(ar_idx, new_ar):
            revised_text[idx] = text

    # Halaman 6-20: Latin lama belum washal dan belum memakai mad;
    # dibangun ulang satu-satu dari Arab berharakat.
    for page in range(6, 21):
        la_idx = index_map[page]["la"]
        assert len(la_idx) == len(revised_arabic[page]), page
        for idx, arabic in zip(la_idx, revised_arabic[page]):
            revised_text[idx] = transliterate_paragraph(arabic)

    # Halaman 61-83: Latin pengguna sudah Washal Taqdiriyyah. Pertahankan
    # bacaannya; bersihkan fragmen aksara asing/harakat liar. Dua baris
    # glyph hitungan Halaman 74 dan penanda wafaq dipertahankan persis.
    for page in [61, 62, 63, 72, 73, 74, 76, 77, 78, 79, 80, 81, 82, 83]:
        for local_idx, idx in enumerate(index_map[page]["la"]):
            old = doc.paragraphs[idx].text
            if page == 74 and local_idx in {2, 3}:
                continue
            new = clean_existing_latin(old)
            if new != old:
                revised_text[idx] = new

    # Halaman 120: dua paragraf Latin lama memuat penggabungan 4+2
    # paragraf Arab. Susunan paragraf dipertahankan; glyph huruf lepas
    # pada daftar panjang tetap ditulis Arab sesuai kaidah hijaiyah.
    group1 = " ".join(revised_arabic[120][:4])
    group2 = " ".join(revised_arabic[120][4:])
    la120 = index_map[120]["la"]
    revised_text[la120[0]] = transliterate_paragraph(group1, preserve_single=True)
    revised_text[la120[1]] = transliterate_paragraph(group2)

    assert problems_total == 0
    return revised_text, revised_arabic


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--report", action="store_true", help="audit harakat tanpa menulis DOCX")
    ap.add_argument(
        "--latin-preview",
        action="store_true",
        help="tampilkan Latin hasil bangun ulang tanpa menulis DOCX",
    )
    args = ap.parse_args()

    before_media = media_hashes(TARGET)
    doc = Document(TARGET)
    before_texts = [p.text for p in doc.paragraphs]
    before_images = image_indices(doc)
    index_map = target_index_map(doc)
    canon = canonical_arabic()
    vocalizer = Vocalizer(doc, canon)

    if args.report:
        all_problems: Counter[str] = Counter()
        total_before = total_after = 0
        total_letters = 0
        for page in TARGET_PAGES:
            texts = [doc.paragraphs[i].text for i in index_map[page]["ar"]]
            revised, problems = vocalizer.vocalize_page(page, texts)
            for old, _new, _q in problems:
                all_problems[old] += 1
            for old, new in zip(texts, revised):
                assert strip_marks(old) == strip_marks(new), (page, old, new)
                total_before += mark_count(old)
                total_after += mark_count(new)
                total_letters += letter_count(old)
            ratio = mark_count("".join(revised)) / max(1, letter_count("".join(revised)))
            print(f"{page}: problems={len(problems)} ratio={ratio:.3f}")
        print(
            f"TOTAL letters={total_letters} marks {total_before}->{total_after} "
            f"ratio={total_after / total_letters:.3f}"
        )
        print(
            f"PROBLEM occurrences={sum(all_problems.values())} "
            f"unique={len(all_problems)}"
        )
        for tok, n in all_problems.most_common(260):
            print(f"{n:3} {tok}")
        return 0

    revised_text, revised_arabic = build_revisions(doc, index_map, vocalizer)
    target_indices = {
        idx
        for page in TARGET_PAGES
        for idx in index_map[page]["ar"] + index_map[page]["la"]
    }
    assert set(revised_text) <= target_indices

    if args.latin_preview:
        for page in list(range(6, 21)) + [120]:
            print(f"\n### {page}")
            for idx in index_map[page]["la"]:
                print(revised_text.get(idx, before_texts[idx]))
        return 0

    for idx, text in revised_text.items():
        set_para_text(doc.paragraphs[idx], text)

    after_texts = [p.text for p in doc.paragraphs]
    assert len(after_texts) == len(before_texts)
    for idx, (old, new) in enumerate(zip(before_texts, after_texts)):
        if idx not in target_indices:
            assert old == new, idx
    assert image_indices(doc) == before_images

    doc.save(TARGET)

    # Verifikasi dari berkas yang benar-benar tersimpan.
    check = Document(TARGET)
    check_texts = [p.text for p in check.paragraphs]
    assert len(check_texts) == len(before_texts)
    for page in TARGET_PAGES:
        got = [check_texts[i] for i in index_map[page]["ar"]]
        assert got == revised_arabic[page], page
        ratio = mark_count("".join(got)) / max(1, letter_count("".join(got)))
        assert ratio >= 0.65, (page, ratio)
    for idx, expected in revised_text.items():
        assert check_texts[idx] == expected, idx
    for idx, (old, new) in enumerate(zip(before_texts, check_texts)):
        if idx not in target_indices:
            assert old == new, idx
    assert image_indices(check) == before_images
    assert media_hashes(TARGET) == before_media

    changed = sum(before_texts[i] != check_texts[i] for i in target_indices)
    total_before = sum(
        mark_count(before_texts[i])
        for page in TARGET_PAGES
        for i in index_map[page]["ar"]
    )
    total_after = sum(
        mark_count(check_texts[i])
        for page in TARGET_PAGES
        for i in index_map[page]["ar"]
    )
    total_letters = sum(
        letter_count(check_texts[i])
        for page in TARGET_PAGES
        for i in index_map[page]["ar"]
    )
    digest = hashlib.sha256(TARGET.read_bytes()).hexdigest()
    print(
        f"AUDIT30_HARAKAT_WASHAL_PASS pages=30 changed_paragraphs={changed} "
        f"marks={total_before}->{total_after} ratio={total_after / total_letters:.3f} "
        f"paragraphs={len(check_texts)} images={len(before_images)} "
        f"bytes={TARGET.stat().st_size} sha256={digest}"
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
