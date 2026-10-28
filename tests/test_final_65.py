"""Structural regression tests; not a claim of scholarly/source verification."""
import re,unittest,zipfile
from pathlib import Path
from docx import Document
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[1]
class MasterTests(unittest.TestCase):
    def test_layers(self):
        md=(ROOT/'REVISI_FINAL_5_LAPIS_BERSIH_HAL01-65.md').read_text()
        blocks=list(filter(str.strip,re.split(r'(?=^## HALAMAN )',md,flags=re.M)))
        self.assertEqual([int(re.match(r'## HALAMAN (\d+)',b).group(1)) for b in blocks],list(range(1,66)))
        expected=['arab','latin','Terjemahan','Syarah dan Penjelasan','Khususiyah dan Tatacara Praktik']
        for b in blocks:
            self.assertEqual(re.findall(r'^### ([^\n]+)$',b,re.M),expected)
            for m in re.finditer(r'^### ([^\n]+)\n(.*?)(?=^### |\Z)',b,re.M|re.S):
                self.assertGreater(len(m.group(2).strip()),15)
    def test_docx(self):
        path=ROOT/'REVISI_FINAL_5_LAPIS_BERSIH_HAL01-65.docx'
        d=Document(path)
        self.assertEqual(sum(p.style.name=='Heading 2' for p in d.paragraphs),65)
        self.assertEqual(sum(p.style.name=='Heading 3' for p in d.paragraphs),325)
        self.assertEqual(len(d.tables),17) # contents + 16 wafaq
        for t in d.tables[1:]:
            self.assertEqual(len(t._tbl.tblPr.findall(qn('w:bidiVisual'))),1)
        for p in d.paragraphs:
            self.assertNotIn('`',p.text);self.assertNotIn('*',p.text)
        with zipfile.ZipFile(path) as z:
            self.assertEqual(len([n for n in z.namelist() if n.startswith('word/media/')]),16)
        self.assertLess(path.stat().st_size,2500000)
    def test_preview(self):
        s=(ROOT/'PREVIEW_INDEX_HAL01-65.html').read_text()
        self.assertEqual(len(re.findall(r'<article id="hal',s)),65)
        self.assertEqual(s.count('<section dir='),325)
        self.assertEqual(s.count('<img '),16)
if __name__=='__main__':unittest.main()
