"""Structural/compliance tests for the Al-Qorin 123-unit standard pilot.

These tests prove file structure, logged deletions, crop identity, DOCX XML
direction/alignment, and recoverability of the embedded Amiri fonts. They do
not by themselves prove semantic completion of all 123 units.
"""
import hashlib
import json
import re
import unittest
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from docx import Document
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
ALQ = ROOT / "al_qorin"
FULL_MD = ALQ / "REVISI_AL_QORIN_123_KERJA.md"
PILOT_MD = ALQ / "REVISI_AL_QORIN_PILOT.md"
PILOT_DOCX = ALQ / "AL_QORIN_PILOT_5_LAPIS.docx"
MANIFEST = ALQ / "MANIFEST_WAFAQ_AL_QORIN.md"
NO_CALC = "«tidak ada hitungan di halaman ini»"
LAYERS = [
    "1. Arab",
    "2. Latin",
    "3. Terjemahan",
    "4. Metode, Penentuan Jam, dan Catatan Penyunting",
    "5. Khususiyah — Waktu : Media : Bukhur",
]
SUBS = ["1. Metode", "2. Penentuan Jam", "Catatan penyunting"]
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
REL = "{http://schemas.openxmlformats.org/package/2006/relationships}"
CT = "{http://schemas.openxmlformats.org/package/2006/content-types}"


def unit_blocks(md):
    blocks = []
    for block in re.split(r"(?=^## )", md, flags=re.M):
        m = re.match(r"^##\s+(.+)$", block, flags=re.M)
        if m:
            heading = m.group(1).strip()
            uid = "SAMPUL" if heading == "SAMPUL" else heading.split()[0]
            blocks.append((uid, heading, block))
    return blocks


def deobfuscate(data, guid):
    key = bytes.fromhex(guid.strip("{}").replace("-", ""))
    key32 = key[::-1] + key[::-1]
    return bytes(b ^ key32[i] for i, b in enumerate(data[:32])) + data[32:]


class MarkdownStructureTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.full = FULL_MD.read_text(encoding="utf-8")
        cls.pilot = PILOT_MD.read_text(encoding="utf-8")
        cls.full_blocks = unit_blocks(cls.full)
        cls.pilot_blocks = unit_blocks(cls.pilot)

    def assert_block_shape(self, blocks):
        for uid, _heading, block in blocks:
            self.assertEqual(re.findall(r"^### ([^\n]+)$", block, re.M), LAYERS, uid)
            self.assertEqual(re.findall(r"^#### ([^\n]+)$", block, re.M), SUBS, uid)
            self.assertEqual(block.count(NO_CALC), 1, uid)
            for layer in LAYERS:
                m = re.search(
                    rf"^### {re.escape(layer)}\n(.*?)(?=^### |\Z)", block, flags=re.M | re.S
                )
                self.assertIsNotNone(m, (uid, layer))
                self.assertGreater(len(m.group(1).strip()), 10, (uid, layer))

    def test_full_123_units_in_order(self):
        expected = ["SAMPUL"] + [f"M{i:02d}" for i in range(1, 123)]
        self.assertEqual([u for u, _h, _b in self.full_blocks], expected)
        self.assert_block_shape(self.full_blocks)

    def test_full_logged_deletions_and_preserved_labels(self):
        self.assertFalse(any(line.strip().startswith("Qarinah ") for line in self.full.splitlines()))
        self.assertEqual(self.full.count("[CATATAN PENERJEMAH"), 25)
        log = (ALQ / "LOG_RESTRUKTUR_AWAL.md").read_text(encoding="utf-8")
        self.assertIn("= 569", log)
        self.assertIn("Kalimat Qarinah sebaris", log)

    def test_pilot_units_and_crops(self):
        expected = ["SAMPUL", "M01", "M02", "M03", "M04", "M05", "M58", "M65", "M66", "M67", "M69", "M122"]
        self.assertEqual([u for u, _h, _b in self.pilot_blocks], expected)
        self.assert_block_shape(self.pilot_blocks)
        refs = [int(x) for x in re.findall(r"wafaq_qorin/(\d+)\.jpg", self.pilot)]
        self.assertEqual(refs, [58, 65, 66, 67, 69])
        for page in refs:
            self.assertTrue((ALQ / f"wafaq_qorin/{page}.jpg").exists())

    def test_manifest_hashes_and_insertion_set(self):
        rows = []
        for line in MANIFEST.read_text(encoding="utf-8").splitlines():
            if re.match(r"^\|\s*\d+\s*\|", line):
                cells = [c.strip() for c in line.strip().strip("|").split("|")]
                rows.append(cells)
        self.assertEqual(len(rows), 32)
        insert_pages = []
        for page, _unit, path, _dim, _size, insert, _cap, sha in rows:
            p = ROOT / path.strip("`")
            self.assertTrue(p.exists(), path)
            self.assertTrue(hashlib.sha256(p.read_bytes()).hexdigest().startswith(sha.strip("`")))
            if insert == "YA":
                insert_pages.append(int(page))
        self.assertEqual(len(insert_pages), 25)
        refs = [int(x) for x in re.findall(r"wafaq_qorin/(\d+)\.jpg", self.full)]
        self.assertEqual(refs, insert_pages)

    def test_pilot_source_layers_unchanged_when_work_data_present(self):
        revised_path = ALQ / "work/revised_units.json"
        pilot_path = ALQ / "work/pilot_units.json"
        if not (revised_path.exists() and pilot_path.exists()):
            self.skipTest("ignored work JSON not present in this checkout")
        revised = {u["unit_id"]: u for u in json.loads(revised_path.read_text(encoding="utf-8"))}
        pilot = json.loads(pilot_path.read_text(encoding="utf-8"))
        for u in pilot:
            src = revised[u["unit_id"]]
            for key in ("arab", "latin", "terjemahan", "metode", "rajah_blocks", "jam"):
                self.assertEqual(u[key], src[key], (u["unit_id"], key))


class FullWorkingDocxTests(unittest.TestCase):
    """The 123-unit DOCX is a mechanical working build, not a semantic final."""

    @classmethod
    def setUpClass(cls):
        cls.path = ALQ / "AL_QORIN_123_KERJA_5_LAPIS.docx"
        if not cls.path.exists():
            raise unittest.SkipTest("full working DOCX has not been built")
        cls.doc = Document(str(cls.path))
        with zipfile.ZipFile(str(cls.path)) as z:
            cls.names = z.namelist()

    def test_full_working_docx_counts(self):
        self.assertEqual(sum(p.style.name == "Heading 1" for p in self.doc.paragraphs), 123)
        self.assertEqual(sum(p.style.name == "Heading 2" for p in self.doc.paragraphs), 615)
        self.assertEqual(sum(p.style.name == "Heading 3" for p in self.doc.paragraphs), 369)
        self.assertEqual(len(self.doc.tables), 4)
        self.assertEqual(len([n for n in self.names if n.startswith("word/media/")]), 25)
        self.assertEqual(len([n for n in self.names if n.startswith("word/fonts/")]), 4)
        texts = [p.text for p in self.doc.paragraphs]
        self.assertEqual(sum(t.count(NO_CALC) for t in texts), 123)
        self.assertEqual(sum(t.count("[CATATAN PENERJEMAH") for t in texts), 25)
        self.assertFalse(any(t.strip().startswith("Qarinah ") for t in texts))


class PilotDocxTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.doc = Document(str(PILOT_DOCX))
        with zipfile.ZipFile(str(PILOT_DOCX)) as z:
            cls.parts = {name: z.read(name) for name in z.namelist()}

    @staticmethod
    def pval(p, tag):
        if p._p.pPr is None:
            return None
        el = p._p.pPr.find(qn(tag))
        return el.get(qn("w:val")) if el is not None else None

    def test_heading_table_and_text_counts(self):
        self.assertEqual(sum(p.style.name == "Heading 1" for p in self.doc.paragraphs), 12)
        self.assertEqual(sum(p.style.name == "Heading 2" for p in self.doc.paragraphs), 60)
        self.assertEqual(sum(p.style.name == "Heading 3" for p in self.doc.paragraphs), 36)
        self.assertEqual([(len(t.rows), len(t.columns)) for t in self.doc.tables],
                         [(1, 2), (4, 3), (2, 10), (5, 4)])
        all_text = [p.text for p in self.doc.paragraphs]
        all_text += [c.text for t in self.doc.tables for r in t.rows for c in r.cells]
        self.assertEqual(sum(t.count(NO_CALC) for t in all_text), 12)
        self.assertFalse(any(t.strip().startswith("Qarinah ") for t in all_text))
        self.assertFalse(any("*" in t or "`" in t for t in all_text))

    def test_arabic_and_latin_direction_flags(self):
        ctx = None
        arab = latin = 0
        for p in self.doc.paragraphs:
            if p.style.name == "Heading 1":
                ctx = None
                continue
            if p.style.name == "Heading 2":
                ctx = p.text.strip()
                continue
            if not p.text.strip():
                continue
            if ctx == "1. Arab":
                arab += 1
                self.assertEqual((self.pval(p, "w:bidi"), self.pval(p, "w:jc")), ("1", "start"), p.text[:80])
            elif ctx == "2. Latin":
                latin += 1
                self.assertEqual((self.pval(p, "w:bidi"), self.pval(p, "w:jc")), ("0", "left"), p.text[:80])
        self.assertGreater(arab, 50)
        self.assertGreater(latin, 40)

    def test_wafaq_images_centered_and_byte_identical(self):
        image_paras = [p for p in self.doc.paragraphs if p._p.xpath('.//*[local-name()="blip"]')]
        self.assertEqual(len(image_paras), 5)
        for p in image_paras:
            self.assertEqual(self.pval(p, "w:jc"), "center")
        media = [self.parts[f"word/media/image{i}.jpg"] for i in range(1, 6)]
        crops = [(ALQ / f"wafaq_qorin/{p}.jpg").read_bytes() for p in (58, 65, 66, 67, 69)]
        self.assertEqual([hashlib.sha256(x).hexdigest() for x in media],
                         [hashlib.sha256(x).hexdigest() for x in crops])

    def test_tables_centered_without_bidi_visual(self):
        for table in self.doc.tables:
            tbl_jc = table._tbl.tblPr.find(qn("w:jc"))
            self.assertIsNotNone(tbl_jc)
            self.assertEqual(tbl_jc.get(qn("w:val")), "center")
            self.assertEqual(len(table._tbl.tblPr.findall(qn("w:bidiVisual"))), 0)
            for row in table.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        self.assertEqual(self.pval(p, "w:jc"), "center")

    def test_amiri_fonts_declared_and_embedded_recoverably(self):
        for style in ("Normal", "Title", "Heading 1", "Heading 2", "Heading 3"):
            self.assertEqual(self.doc.styles[style].font.name, "Amiri", style)
        settings = ET.fromstring(self.parts["word/settings.xml"])
        self.assertIsNotNone(settings.find(f"{W}embedTrueTypeFonts"))
        self.assertIsNone(settings.find(f"{W}saveSubsetFonts"))
        ctypes = ET.fromstring(self.parts["[Content_Types].xml"])
        self.assertTrue(any(el.get("Extension") == "odttf" for el in ctypes.findall(f"{CT}Default")))
        font_table = ET.fromstring(self.parts["word/fontTable.xml"])
        Amiri = None
        for el in font_table.findall(f"{W}font"):
            if el.get(f"{W}name") == "Amiri":
                Amiri = el
                break
        self.assertIsNotNone(Amiri)
        rels = ET.fromstring(self.parts["word/_rels/fontTable.xml.rels"])
        relmap = {el.get("Id"): el.get("Target") for el in rels.findall(f"{REL}Relationship")}
        expected = {
            "embedRegular": "Amiri-Regular.ttf",
            "embedBold": "Amiri-Bold.ttf",
            "embedItalic": "Amiri-Italic.ttf",
            "embedBoldItalic": "Amiri-BoldItalic.ttf",
        }
        for tag, ttf in expected.items():
            emb = Amiri.find(f"{W}{tag}")
            self.assertIsNotNone(emb, tag)
            part = "word/" + relmap[emb.get(f"{R}id")]
            restored = deobfuscate(self.parts[part], emb.get(f"{W}fontKey"))
            self.assertEqual(restored, (ROOT / "fonts" / ttf).read_bytes(), tag)


if __name__ == "__main__":
    unittest.main()
